# -*- coding: utf-8 -*-
"""三环集团（300408.SZ / 06951.HK）详细财务与估值模型 v2（单位：百万元人民币）"""
import sys, json
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from engine import *
from hist_detail import DATA

SRC = 'src/e66ff817-____/'
OUT = sys.argv[1]
wb = openpyxl.Workbook(); wb.remove(wb.active)

def D(*path):
    """DATA lookup returning None if missing."""
    d = DATA
    for p in path:
        if not isinstance(d, dict) or p not in d:
            return None
        d = d[p]
    return d

def DH(*path):
    """history dict keyed by model years ('2023A') from DATA keyed by 'FY2023'."""
    d = D(*path)
    if not d:
        return None
    return {y: d[FY[y]] for y in HIST if FY[y] in d and d[FY[y]] is not None}

# ============================================================ raw Wind sheets
RAW_COLS = ['2019年报', '2020年报', '2021年报', '2022年报', '2023年报', '2024年报', '2025年报', '2025中报', '2026中报']
RAW_HDR = ['2019A', '2020A', '2021A', '2022A', '2023A', '2024A', '2025A', '2025H1', '2026H1']
RAW_MAP = {'2021A': 'E', '2022A': 'F', '2023A': 'G', '2024A': 'H', '2025A': 'I', '2025H1': 'J', '2026H1': 'K'}
raw_rows = {}
SKIP = ('显示币种', '原始币种', '转换汇率', '汇率类型', '税率', '审计意见(境内)', '公告日期', '数据来源', '报表类型')

def build_raw(sheet_name, fname, ttl):
    ws = wb.create_sheet(sheet_name)
    title(ws, ttl, '资料来源：Wind（公司公告值）；原始单位万元，已除以100换算为百万元。蓝字为历史录入值。')
    header_row(ws, 4, ['项目（百万元）', ''] + RAW_HDR)
    src = openpyxl.load_workbook(SRC + fname, data_only=True).active
    rows = list(src.iter_rows(values_only=True)); hdr = rows[1]
    idx = [hdr.index(c) for c in RAW_COLS]
    r = 5
    for row in rows[2:]:
        lab = row[0]
        if lab is None or str(lab).strip() in SKIP:
            continue
        vals = [row[i] for i in idx]
        if all(v is None for v in vals):
            continue
        s = str(lab)
        put(ws, f'A{r}', s.rstrip(), bold=not s.startswith(' '))
        for j, v in enumerate(vals):
            if isinstance(v, (int, float)):
                ps = '每股' in s
                put(ws, f'{CL(3 + j)}{r}', round(v if ps else v / 100, 6), PS if ps else NUM)
        raw_rows[(sheet_name, s.strip())] = r
        r += 1
    ws.column_dimensions['A'].width = 46
    for j in range(len(RAW_HDR)):
        ws.column_dimensions[CL(3 + j)].width = 12
    ws.freeze_panes = 'C5'; ws.sheet_view.showGridLines = False

HL, HB, HCF = '历史利润表', '历史资产负债表', '历史现金流量表'

def RAW(sheet, label, year):
    return f"'{sheet}'!{RAW_MAP[year]}{raw_rows[(sheet, label)]}"

def RAWS(sheet, labels, year, neg=False):
    parts = [f'N({RAW(sheet, l, year)})' for l in labels if (sheet, l) in raw_rows]
    return ('=-(' if neg else '=') + '+'.join(parts) + (')' if neg else '')

# ============================================================ sheet objects (layout first)
S_REV, S_COST, S_OPEX, S_CAP, S_WC, S_FIN, S_TAX, S_M, S_SH, S_RET = (
    '收入拆解', '成本拆解', '人员与费用', '资本开支与折旧', '营运资金', '融资与资金', '税务与其他损益', '三表预测', '股本与每股', '回报分析')

def A(key, col):      # per-year active assumption
    return f"'假设'!{col}{REG[('假设', key)]}"

def A1(key):          # single assumption (absolute)
    return f"'假设'!$C${REG[('假设1', key)]}"

def G(sheet, key, col):
    return X(sheet, key, col)

def L(key, col):      # same-sheet reference
    return X(CUR[0], key, col)

def hist_formula(fn):
    return lambda y: fn(YC[y], y)

H23 = ['2023A', '2024A', '2025A']

# ---------------------------------------------------------------- 收入拆解
rv = Sheet(S_REV, '收入拆解：量价、分部与SOFC隔膜片子模型',
           '分部口径：2024A-2025A为2025年年报新四分部口径（与H股招股章程一致）；量价拆分为假设驱动（公司未按分部披露销量）。')
rv.sec('A. 公司口径量价（年报“产销量”，万只/万片÷1万=亿只/片）')
vol = D('volume')
rv.row('vol_sales', '销售量', '亿只/片', hist={y: v / 1e4 for y, v in (DH('volume', '销售量') or {}).items()}, fmt=NUM, src='年报“产销量情况分析”（单位万只/万片）')
rv.row('vol_prod', '生产量', '亿只/片', hist={y: v / 1e4 for y, v in (DH('volume', '生产量') or {}).items()}, fmt=NUM)
rv.row('vol_inv', '库存量', '亿只/片', hist={y: v / 1e4 for y, v in (DH('volume', '库存量') or {}).items()}, fmt=NUM)
rv.row('asp', '综合单价（营业收入÷销售量）', '元/千只', hist=lambda y: (f"=IFERROR({G(S_M,'rev',YC[y])}*10/{L('vol_sales',YC[y])},0)" if DH('volume','销售量') and y in DH('volume','销售量') else None), fmt=NUM2,
       src='综合单价受产品结构影响大（MLCC小尺寸占量大、单价低），仅作趋势参考')
rv.row('vol_g', '销售量同比', '%', hist=lambda y: (f"=IFERROR({L('vol_sales',YC[y])}/{L('vol_sales',prev(YC[y]))}-1,0)" if DH('volume','销售量') and y in DH('volume','销售量') and FY[y] != min(D('volume','销售量')) else None), fmt=PCT)
rv.row('asp_g', '综合单价同比', '%', hist=lambda y: (f"=IFERROR({L('asp',YC[y])}/{L('asp',prev(YC[y]))}-1,0)" if DH('volume','销售量') and y in DH('volume','销售量') and FY[y] != min(D('volume','销售量')) else None), fmt=PCT)

def seg_vp(tag, name, h24, h25, vkey, pkey, src):
    rv.sec(f'{name}')
    rv.row(f'{tag}_v', '  销量增速（假设）', '%', fcst=lambda c: f"={A(vkey,c)}", fmt=PCT)
    rv.row(f'{tag}_p', '  单价变动（假设）', '%', fcst=lambda c: f"={A(pkey,c)}", fmt=PCT)
    rv.row(f'{tag}_idx_v', '  销量指数（2025A=100）', '', hist={'2025A': 100}, fcst=lambda c: f"={L(tag+'_idx_v',prev(c))}*(1+{L(tag+'_v',c)})", fmt=NUM)
    rv.row(f'{tag}_idx_p', '  单价指数（2025A=100）', '', hist={'2025A': 100}, fcst=lambda c: f"={L(tag+'_idx_p',prev(c))}*(1+{L(tag+'_p',c)})", fmt=NUM)
    rv.row(f'{tag}_rev', f'{name}收入', '百万元', hist={'2024A': h24, '2025A': h25},
           fcst=lambda c: f"={L(tag+'_rev',prev(c))}*(1+{L(tag+'_v',c)})*(1+{L(tag+'_p',c)})", bold=True, src=src)
    rv.row(f'{tag}_cv', '  其中：量的贡献', '百万元', fcst=lambda c: f"={L(tag+'_rev',prev(c))}*{L(tag+'_v',c)}")
    rv.row(f'{tag}_cp', '  其中：价的贡献', '百万元', fcst=lambda c: f"={L(tag+'_rev',prev(c))}*(1+{L(tag+'_v',c)})*{L(tag+'_p',c)}")
    rv.row(f'{tag}_g', '  同比', '%', hist={'2025A': '=IFERROR(G{'+tag+'_rev}/F{'+tag+'_rev}-1,0)'}, fcst=lambda c: f"=IFERROR({L(tag+'_rev',c)}/{L(tag+'_rev',prev(c))}-1,0)", fmt=PCT)

seg_vp('ec', '电子元件（MLCC、陶瓷封装基座、电阻等）', 2298.159034, 3308.089586, 'ec_vol', 'ec_px', '2025年年报：电子元件3,308,089,586.22元（2024年2,298,159,034.08元）')
seg_vp('cd', '通信器件（陶瓷插芯、套筒等光通信产品）', 2305.984656, 2594.309892, 'cd_vol', 'cd_px', '2025年年报：通信器件2,594,309,892.31元（2024年2,305,984,655.76元）')

rv.sec('电子及陶瓷材料（陶瓷基片/基板、电子浆料、SOFC隔膜片等）')
rv.row('gw', '  Bloom电池片层需求', 'GW', hist=lambda y: f"={A('sofc_gw','G')}" if y == '2025A' else None, fcst=lambda c: f"={A('sofc_gw',c)}", fmt=NUM2, src='见“假设”SOFC子模型')
rv.row('w', '  单片功率', 'W/片', fcst=lambda c: f"={A('sofc_w',c)}", fmt='0.0')
rv.row('sheets_gw', '  每GW隔膜片需求', '百万片/GW', fcst=lambda c: f"=IFERROR(1000/{L('w',c)},0)", fmt=NUM)
rv.row('demand', '  Bloom隔膜片需求', '百万片', fcst=lambda c: f"={L('gw',c)}*{L('sheets_gw',c)}", fmt=NUM)
rv.row('share', '  三环份额', '%', fcst=lambda c: f"={A('sofc_sh',c)}", fmt=PCT)
rv.row('svol', '  三环隔膜片出货', '百万片', fcst=lambda c: f"={L('demand',c)}*{L('share',c)}", fmt=NUM)
rv.row('spx', '  单价', '元/片', fcst=lambda c: f"={A('sofc_px',c)}", fmt=NUM2)
rv.row('sofc_rev', '  SOFC隔膜片收入', '百万元', hist={'2023A': 449.318728, '2024A': 323.869648, '2025A': 330.717999},
       fcst=lambda c: f"={L('svol',c)}*{L('spx',c)}", bold=True,
       src='历史为年报前五大客户“客户1”销售额（2023A 449,318,727.96元；2024A 323,869,648.00元；2025A 330,717,999.78元），市场普遍对应Bloom')
rv.row('sofc_valgw', '  每GW隔膜片价值量（100%份额）', '百万元/GW', fcst=lambda c: f"={L('sheets_gw',c)}*{L('spx',c)}")
rv.row('cmx_rev', '  其他电子及陶瓷材料（不含SOFC）', '百万元', hist={'2024A': '=F{cm_rev}-F{sofc_rev}', '2025A': '=G{cm_rev}-G{sofc_rev}'},
       fcst=lambda c: f"={L('cmx_rev',prev(c))}*(1+{A('g_cm',c)})")
rv.row('cm_rev', '电子及陶瓷材料收入', '百万元', hist={'2024A': 1677.232781, '2025A': 1959.378017},
       fcst=lambda c: f"={L('cmx_rev',c)}+{L('sofc_rev',c)}", bold=True, src='2025年年报：电子及陶瓷材料1,959,378,016.92元（2024年1,677,232,780.51元）')

rv.sec('设备组件、其他主营与其他业务')
rv.row('eq_rev', '设备组件收入', '百万元', hist={'2024A': 536.824730, '2025A': 601.753410}, fcst=lambda c: f"={L('eq_rev',prev(c))}*(1+{A('g_eq',c)})", bold=True, src='2025年年报：601,753,409.61元（2024年536,824,729.50元）')
rv.row('ot_rev', '其他主营收入', '百万元', hist={'2024A': 447.931467, '2025A': 404.996373}, fcst=lambda c: f"={L('ot_rev',prev(c))}*(1+{A('g_ot',c)})", bold=True, src='2025年年报：404,996,373.28元（2024年447,931,467.25元）')
rv.row('ob_rev', '其他业务收入', '百万元', hist={'2024A': 108.858154, '2025A': 138.626616}, fcst=lambda c: f"={A('rev_ob',c)}", bold=True, src='2025年年报：138,626,615.77元；2026H1半年报“其他”342.6为其他主营+其他业务合计（新口径下设备组件并入“电子、通信元件及材料”）')

rv.sec('合计与结构')
rv.row('rev_tot', '营业收入合计', '百万元', hist={'2024A': '=F{ec_rev}+F{cd_rev}+F{cm_rev}+F{eq_rev}+F{ot_rev}+F{ob_rev}', '2025A': '=G{ec_rev}+G{cd_rev}+G{cm_rev}+G{eq_rev}+G{ot_rev}+G{ob_rev}'},
       fcst=lambda c: f"={L('ec_rev',c)}+{L('cd_rev',c)}+{L('cm_rev',c)}+{L('eq_rev',c)}+{L('ot_rev',c)}+{L('ob_rev',c)}", bold=True)
rv.row('rev_g', '  同比', '%', hist={'2025A': '=IFERROR(G{rev_tot}/F{rev_tot}-1,0)'}, fcst=lambda c: f"=IFERROR({L('rev_tot',c)}/{L('rev_tot',prev(c))}-1,0)", fmt=PCT)
for tag, nm in (('ec', '电子元件'), ('cd', '通信器件'), ('cm', '电子及陶瓷材料'), ('sofc', '  其中SOFC隔膜片'), ('eq', '设备组件')):
    rv.row(f'{tag}_mix', f'  {nm}占比', '%', hist={y: f"=IFERROR({YC[y]}{{{tag}_rev}}/{YC[y]}{{rev_tot}},0)" for y in ('2024A', '2025A')},
           fcst=lambda c, t=tag: f"=IFERROR({L(t+'_rev',c)}/{L('rev_tot',c)},0)", fmt=PCT)
rv.sec('收入增量拆解（同比增加额，百万元）')
for tag, nm in (('ec', '电子元件'), ('cd', '通信器件'), ('cm', '电子及陶瓷材料'), ('eq', '设备组件'), ('ot', '其他主营'), ('ob', '其他业务')):
    rv.row(f'{tag}_d', f'  {nm}', '百万元', hist={'2025A': f"=G{{{tag}_rev}}-F{{{tag}_rev}}"}, fcst=lambda c, t=tag: f"={L(t+'_rev',c)}-{L(t+'_rev',prev(c))}")
rv.row('d_tot', '  合计', '百万元', hist={'2025A': '=SUM(G{ec_d}:G{ob_d})'}, fcst=lambda c: f"=SUM({L('ec_d',c)}:{L('ob_d',c)})", bold=True)
rv.sec('地区结构')
rv.row('os_share', '境外收入占比', '%', hist={y: v for y, v in (DH('region', 'os_share') or {}).items()}, fcst=lambda c: f"={A('os_share',c)}", fmt=PCT, src='年报“营业收入构成-分地区”；H股招股章程：2023-2025年海外收入占比20.3%/17.9%/17.4%')
rv.row('os_rev', '境外收入', '百万元', fcst=lambda c: f"={L('rev_tot',c)}*{L('os_share',c)}")
rv.row('dom_rev', '境内收入', '百万元', fcst=lambda c: f"={L('rev_tot',c)}-{L('os_rev',c)}")
rv.sec('半年度校验（2026H1半年报新口径：电子、通信元件及材料=电子元件+通信器件+电子及陶瓷材料+设备组件；其他=其他主营+其他业务）')
rv.row('h1_core', '2026H1 电子、通信元件及材料（实际）', '百万元', fcst=lambda c: 6080.743359 if c == 'H' else None, src='2026年半年报：6,080,743,358.96元，同比+54.97%（对应2025H1可比基数约3,923.8）')
rv.row('h1_oth', '2026H1 其他（实际）', '百万元', fcst=lambda c: f"={L('h1',c)}-{L('h1_core',c)}" if c == 'H' else None)
rv.row('h2_core', '2026E隐含下半年 电子、通信元件及材料', '百万元', fcst=lambda c: f"={L('ec_rev',c)}+{L('cd_rev',c)}+{L('cm_rev',c)}+{L('eq_rev',c)}-{L('h1_core',c)}" if c == 'H' else None)
rv.row('h2_oth', '2026E隐含下半年 其他', '百万元', fcst=lambda c: f"={L('ot_rev',c)}+{L('ob_rev',c)}-{L('h1_oth',c)}" if c == 'H' else None,
       src='2025年可比：上半年约225.0、全年543.6（其他主营405.0+其他业务138.6）')
rv.row('h1', '2026H1营业收入（半年报）', '百万元', fcst=lambda c: f"={RAW(HL,'营业收入','2026H1')}" if c == 'H' else None)
rv.row('h2', '2026E隐含下半年收入', '百万元', fcst=lambda c: f"={L('rev_tot',c)}-{L('h1',c)}" if c == 'H' else None)
rv.row('h2h1', '2026E下半年/上半年', 'x', fcst=lambda c: f"=IFERROR({L('h2',c)}/{L('h1',c)},0)" if c == 'H' else None, fmt='0.00"x"',
       src='参照：2025年下半年/上半年约1.17x；2026Q2单季收入3,742百万元')

# ---------------------------------------------------------------- 成本拆解
co = Sheet(S_COST, '成本拆解：分部毛利率与营业成本按性质', '分部毛利率为主驱动（生产人工与生产折旧偏离标定值的部分另行计入）；按性质拆分时人工来自人员模型、折旧来自资本开支模型、其他制造费用按占收入比例假设，原材料为倒挤项。')
co.sec('A. 分部毛利率（采用值，含敏感性冲击）')
SEGS = [('ec', '电子元件', 0.4437, 0.4230), ('cd', '通信器件', 0.4393, 0.4436), ('cmx', '其他电子及陶瓷材料', None, None), ('sofc', 'SOFC隔膜片', None, None),
        ('eq', '设备组件', None, 0.5846), ('ot', '其他主营', None, 0.0911), ('ob', '其他业务', None, 0.7957)]
GMK = {'ec': 'gm_ec', 'cd': 'gm_cd', 'cmx': 'gm_cm', 'sofc': 'sofc_gm', 'eq': 'gm_eq', 'ot': 'gm_ot', 'ob': 'gm_ob'}
for tag, nm, g24, g25 in SEGS:
    h = {}
    if g24 is not None: h['2024A'] = g24
    if g25 is not None: h['2025A'] = g25
    co.row(f'gm_{tag}', f'{nm} 毛利率', '%', hist=h or None, fcst=lambda c, k=GMK[tag]: f"={A(k,c)}", fmt=PCT,
           src='2025年年报分产品毛利率及同比增减倒推；设备组件/其他主营/其他业务按Wind主营构成' if tag in ('ec', 'cd', 'eq') else None)
co.row('gm_cm_all', '电子及陶瓷材料（分部合计）毛利率', '%', hist={'2024A': 0.4057, '2025A': 0.3809},
       fcst=lambda c: f"=IFERROR(1-{L('cogs_cm',c)}/{G(S_REV,'cm_rev',c)},0)", fmt=PCT)
co.sec('B. 分部营业成本')
REVK = {'ec': 'ec_rev', 'cd': 'cd_rev', 'cmx': 'cmx_rev', 'sofc': 'sofc_rev', 'eq': 'eq_rev', 'ot': 'ot_rev', 'ob': 'ob_rev'}
H25COGS = {'ec': 1908.641599, 'cd': 1443.603418, 'eq': 249.9511, 'ot': 368.1160, 'ob': 28.3312}
for tag, nm, *_ in SEGS:
    h = {'2025A': H25COGS[tag]} if tag in H25COGS else None
    co.row(f'cogs_{tag}', f'{nm} 营业成本', '百万元', hist=h, fcst=lambda c, t=tag: f"={G(S_REV,REVK[t],c)}*(1-{L('gm_'+t,c)})")
co.row('cogs_cm', '电子及陶瓷材料 营业成本（合计）', '百万元', hist={'2025A': 1213.118348}, fcst=lambda c: f"={L('cogs_cmx',c)}+{L('cogs_sofc',c)}")
co.row('cogs_seg', '分部营业成本小计（按毛利率假设）', '百万元', hist={'2025A': '=G{cogs_ec}+G{cogs_cd}+G{cogs_cm}+G{cogs_eq}+G{cogs_ot}+G{cogs_ob}'},
       fcst=lambda c: f"={L('cogs_ec',c)}+{L('cogs_cd',c)}+{L('cogs_cm',c)}+{L('cogs_eq',c)}+{L('cogs_ot',c)}+{L('cogs_ob',c)}")
co.row('adj_lab', '加：生产人工偏离标定值', '百万元', fcst=lambda c: f"={L('labor',c)}-{A('lab_ref',c)}",
       src='人数或人均薪酬假设偏离标定值时，差额直接计入营业成本（基准情景为0）')
co.row('adj_da', '加：生产折旧摊销偏离标定值', '百万元', fcst=lambda c: f"={L('oh_da',c)}-{A('da_ref',c)}",
       src='资本开支、折旧率或情景变化引起的生产折旧差额直接计入营业成本（基准情景为0）')
co.row('cogs_tot', '营业成本合计', '百万元', hist=lambda y: f"={RAW(HL,'营业成本',y)}",
       fcst=lambda c: f"={L('cogs_seg',c)}+{L('adj_lab',c)}+{L('adj_da',c)}", bold=True)
co.row('gp_tot', '毛利合计', '百万元', hist=lambda y: f"={G(S_REV,'rev_tot',YC[y])}-{L('cogs_tot',YC[y])}" if y in ('2024A', '2025A') else f"={RAW(HL,'营业收入',y)}-{RAW(HL,'营业成本',y)}",
       fcst=lambda c: f"={G(S_REV,'rev_tot',c)}-{L('cogs_tot',c)}", bold=True)
co.row('gm_tot', '综合毛利率', '%', hist=lambda y: f"=IFERROR({L('gp_tot',YC[y])}/{RAW(HL,'营业收入',y)},0)", fcst=lambda c: f"=IFERROR({L('gp_tot',c)}/{G(S_REV,'rev_tot',c)},0)", fmt=PCT, bold=True)
co.row('gp_sofc', 'SOFC隔膜片毛利', '百万元', fcst=lambda c: f"={G(S_REV,'sofc_rev',c)}-{L('cogs_sofc',c)}")
co.row('gp_sofc_sh', '  SOFC毛利占公司毛利', '%', fcst=lambda c: f"=IFERROR({L('gp_sofc',c)}/{L('gp_tot',c)},0)", fmt=PCT)
co.sec('C. 营业成本按性质（人工来自人员模型、生产折旧来自资本开支模型、其他制造费用按假设，原材料为倒挤项）')
cn = D('cogs_nature') or {}
co.row('mat', '原材料（预测为倒挤：营业成本-人工-生产折旧-其他制造费用）', '百万元', hist=DH('cogs_nature', '原材料'), fcst=lambda c: f"={L('cogs_tot',c)}-{L('labor',c)}-{L('oh',c)}", src='年报“营业成本构成”（工业）')
co.row('labor', '人工费用', '百万元', hist=DH('cogs_nature', '人工费用'), fcst=lambda c: f"={G(S_OPEX,'cogs_staff',c)}")
co.row('oh', '制造费用', '百万元', hist=DH('cogs_nature', '制造费用'), fcst=lambda c: f"={L('oh_da',c)}+{L('oh_oth',c)}")
co.row('oh_da', '  其中：生产用折旧摊销', '百万元', hist=lambda y: (f"={G(S_CAP,'da_tot',YC[y])}*{A1('da_cogs')}" if y in H23 else None), fcst=lambda c: f"={G(S_CAP,'da_cogs',c)}",
       src='历史按折旧摊销总额×生产分摊比例估算')
co.row('oh_oth', '  其中：能源、辅料及其他制造费用', '百万元', hist=lambda y: (f"={L('oh',YC[y])}-{L('oh_da',YC[y])}" if y in H23 else None), fcst=lambda c: f"={G(S_REV,'rev_tot',c)}*{A('oh_oth_r',c)}")
co.row('nat_chk', '  校验：性质合计-营业成本', '百万元', hist=lambda y: (f"=ROUND({L('mat',YC[y])}+{L('labor',YC[y])}+{L('oh',YC[y])}-{L('cogs_tot',YC[y])},1)" if y in H23 else None),
       fcst=lambda c: f"=ROUND({L('mat',c)}+{L('labor',c)}+{L('oh',c)}-{L('cogs_tot',c)},4)")
co.row('mat_sh', '原材料占营业成本（预测为结果，对照历史53.6%-63.6%）', '%', hist=lambda y: (f"=IFERROR({L('mat',YC[y])}/{L('cogs_tot',YC[y])},0)" if y in H23 else None), fcst=lambda c: f"=IFERROR({L('mat',c)}/{L('cogs_tot',c)},0)", fmt=PCT)
co.row('lab_sh', '人工占营业成本', '%', hist=lambda y: (f"=IFERROR({L('labor',YC[y])}/{L('cogs_tot',YC[y])},0)" if y in H23 else None), fcst=lambda c: f"=IFERROR({L('labor',c)}/{L('cogs_tot',c)},0)", fmt=PCT)
co.row('oh_sh', '制造费用占营业成本', '%', hist=lambda y: (f"=IFERROR({L('oh',YC[y])}/{L('cogs_tot',YC[y])},0)" if y in H23 else None), fcst=lambda c: f"=IFERROR({L('oh',c)}/{L('cogs_tot',c)},0)", fmt=PCT)
co.row('oh_oth_r', '其他制造费用占收入', '%', hist=lambda y: (f"=IFERROR({L('oh_oth',YC[y])}/{RAW(HL,'营业收入',y)},0)" if y in H23 else None), fcst=lambda c: f"=IFERROR({L('oh_oth',c)}/{G(S_REV,'rev_tot',c)},0)", fmt=PCT,
       src='预测为假设驱动（见“假设”）')
co.row('mat_rev', '原材料占收入', '%', hist=lambda y: (f"=IFERROR({L('mat',YC[y])}/{RAW(HL,'营业收入',y)},0)" if y in H23 else None), fcst=lambda c: f"=IFERROR({L('mat',c)}/{G(S_REV,'rev_tot',c)},0)", fmt=PCT)

# ---------------------------------------------------------------- 人员与费用
op = Sheet(S_OPEX, '人员与期间费用：人数×人均薪酬驱动', '员工人数与专业构成来自年报“公司员工情况”；职工薪酬总额取“应付职工薪酬”本期增加；费用按性质拆分来自财务报表附注。')
op.sec('A. 员工人数（期末，人）')
HCK = [('prod', '生产人员'), ('sales', '销售人员'), ('tech', '技术人员'), ('fin', '财务人员'), ('admin', '行政人员')]
for k, nm in HCK:
    op.row(f'hc_{k}', nm, '人', hist=DH('hc', nm), fcst=lambda c, k=k: f"={L('hc_'+k,prev(c))}*(1+{A('hc_'+k+'_g',c)})", fmt=NUM0)
op.row('hc_tot', '员工合计', '人', hist=lambda y: (f"=SUM({YC[y]}{{hc_prod}}:{YC[y]}{{hc_admin}})" if DH('hc', '生产人员') and y in DH('hc', '生产人员') else None),
       fcst=lambda c: f"=SUM({L('hc_prod',c)}:{L('hc_admin',c)})", fmt=NUM0, bold=True)
op.row('hc_g', '  员工合计同比', '%', hist=lambda y: (f"=IFERROR({YC[y]}{{hc_tot}}/{prev(YC[y])}{{hc_tot}}-1,0)" if DH('hc','生产人员') and y in DH('hc','生产人员') and FY[y] != min(D('hc','生产人员')) else None),
       fcst=lambda c: f"=IFERROR({L('hc_tot',c)}/{L('hc_tot',prev(c))}-1,0)", fmt=PCT)
op.row('rev_per_hc', '  人均营业收入', '万元/人', hist=lambda y: (f"=IFERROR({RAW(HL,'营业收入',y)}*100/AVERAGE({prev(YC[y])}{{hc_tot}},{YC[y]}{{hc_tot}}),0)" if DH('hc','生产人员') and y in DH('hc','生产人员') and FY[y] != min(D('hc','生产人员')) else None),
       fcst=lambda c: f"=IFERROR({G(S_REV,'rev_tot',c)}*100/AVERAGE({L('hc_tot',prev(c))},{L('hc_tot',c)}),0)", fmt=NUM)
for k, nm in HCK:
    op.row(f'avg_{k}', f'  {nm}（平均）', '人', hist=lambda y, k=k, nm=nm: (f"=AVERAGE({prev(YC[y])}{{hc_{k}}},{YC[y]}{{hc_{k}}})" if DH('hc',nm) and y in DH('hc',nm) and FY[y] != min(D('hc',nm)) else None),
           fcst=lambda c, k=k: f"=AVERAGE({L('hc_'+k,prev(c))},{L('hc_'+k,c)})", fmt=NUM0)
op.row('avg_tot', '  平均员工合计', '人', hist=lambda y: (f"=SUM({YC[y]}{{avg_prod}}:{YC[y]}{{avg_admin}})" if DH('hc','生产人员') and y in DH('hc','生产人员') and FY[y] != min(D('hc','生产人员')) else None),
       fcst=lambda c: f"=SUM({L('avg_prod',c)}:{L('avg_admin',c)})", fmt=NUM0)
op.sec('B. 职工薪酬总额与人均')
op.row('comp_tot', '职工薪酬总额', '百万元', hist=DH('comp_total'), fcst=lambda c: f"={L('comp_pc',c)}*{L('avg_tot',c)}/100", bold=True, src='应付职工薪酬“本期增加”（短期薪酬+离职后福利）')
op.row('comp_pc', '人均薪酬', '万元/人', hist=lambda y: (f"=IFERROR({YC[y]}{{comp_tot}}*100/{YC[y]}{{avg_tot}},0)" if DH('comp_total') and y in DH('comp_total') and DH('hc','生产人员') and FY[y] != min(D('hc','生产人员')) else None),
       fcst=lambda c: f"={L('comp_pc',prev(c))}*(1+{A('wage_g',c)})", fmt=NUM2)
op.row('comp_g', '  人均薪酬同比', '%', hist={'2025A': '=IFERROR(G{comp_pc}/F{comp_pc}-1,0)'}, fcst=lambda c: f"={A('wage_g',c)}", fmt=PCT)
op.sec('C. 职工薪酬按功能归集（上年×平均总人数变动×(1+人均薪酬增速)；2025年技术人员由3,918人降至1,975人、生产人员增加3,468人，为口径调整）')
def staff_row(key, label, hist, avgkeys):
    avg_now = lambda c: '+'.join(L('avg_' + k, c) for k in avgkeys)
    op.row(key, label, '百万元', hist=hist,
           fcst=lambda c: f"={L(key,prev(c))}*IFERROR(({avg_now(c)})/({avg_now(prev(c))}),1)*(1+{A('wage_g',c)})")
ALLK = ['prod', 'sales', 'tech', 'fin', 'admin']   # 2025年技术人员大幅重分类至生产人员，按功能人数滚动会失真，统一按平均总人数滚动
staff_row('cogs_staff', '计入营业成本的人工', DH('cogs_nature', '人工费用'), ALLK)
staff_row('rd_staff', '研发费用-职工薪酬', DH('opex_nature', '研发费用', '职工薪酬'), ALLK)
staff_row('adm_staff', '管理费用-职工薪酬', DH('opex_nature', '管理费用', '职工薪酬'), ALLK)
staff_row('sell_staff', '销售费用-职工薪酬', DH('opex_nature', '销售费用', '职工薪酬'), ALLK)
op.row('oth_staff', '其他职工薪酬（计入在建工程、存货等，倒挤）', '百万元',
       hist=lambda y: (f"={YC[y]}{{comp_tot}}-{YC[y]}{{cogs_staff}}-{YC[y]}{{rd_staff}}-{YC[y]}{{adm_staff}}-{YC[y]}{{sell_staff}}" if DH('opex_nature','研发费用','职工薪酬') and y in DH('opex_nature','研发费用','职工薪酬') and DH('comp_total') and y in DH('comp_total') else None),
       fcst=lambda c: f"={L('comp_tot',c)}-{L('cogs_staff',c)}-{L('rd_staff',c)}-{L('adm_staff',c)}-{L('sell_staff',c)}",
       src='倒挤项，反映薪酬在成本与费用之间的归集差异及存货中尚未结转的人工')
op.sec('D. 销售费用')
def opex_block(tag, nm, rawlab, staff_key, da_key, others):
    op.row(f'{tag}_staff2', '  职工薪酬', '百万元', hist=lambda y: (f"={YC[y]}{{{staff_key}}}" if D('opex_nature', nm, '职工薪酬') and FY[y] in D('opex_nature', nm, '职工薪酬') else None), fcst=lambda c: f"={L(staff_key,c)}")
    op.row(f'{tag}_da', '  折旧与摊销', '百万元', hist=DH('opex_nature', nm, '折旧摊销'), fcst=lambda c: f"={G(S_CAP,da_key,c)}")
    rows_other = []
    for ok, olab, akey in others:
        op.row(f'{tag}_{ok}', f'  {olab}', '百万元', hist=DH('opex_nature', nm, olab), fcst=lambda c, a=akey: f"={G(S_REV,'rev_tot',c)}*{A(a,c)}")
        rows_other.append(f'{tag}_{ok}')
    op.row(f'{tag}_oth', '  其他', '百万元',
           hist=lambda y: (f"={RAW(HL,rawlab,y)}-{YC[y]}{{{tag}_staff2}}-{YC[y]}{{{tag}_da}}" + ''.join(f"-{YC[y]}{{{r}}}" for r in rows_other) if D('opex_nature', nm, '职工薪酬') and FY[y] in D('opex_nature', nm, '职工薪酬') else None),
           fcst=lambda c: f"={G(S_REV,'rev_tot',c)}*{A(tag+'_oth_r',c)}")
    op.row(f'{tag}_tot', f'{nm}合计', '百万元', hist=lambda y: f"={RAW(HL,rawlab,y)}",
           fcst=lambda c: f"={L(tag+'_staff2',c)}+{L(tag+'_da',c)}" + ''.join(f"+{L(r,c)}" for r in rows_other) + f"+{L(tag+'_oth',c)}", bold=True)
    op.row(f'{tag}_r', f'  {nm}率', '%', hist=lambda y: f"=IFERROR({YC[y]}{{{tag}_tot}}/{RAW(HL,'营业收入',y)},0)", fcst=lambda c: f"=IFERROR({L(tag+'_tot',c)}/{G(S_REV,'rev_tot',c)},0)", fmt=PCT)
opex_block('sell', '销售费用', '销售费用', 'sell_staff', 'da_sell', [])
op.sec('E. 管理费用')
opex_block('adm', '管理费用', '管理费用', 'adm_staff', 'da_adm', [])
op.sec('F. 研发费用')
opex_block('rd', '研发费用', '研发费用', 'rd_staff', 'da_rd', [('mat', '直接投入/材料', 'rd_mat_r')])
op.sec('G. 期间费用合计')
op.row('opex_tot', '销售+管理+研发', '百万元', hist=lambda y: f"={YC[y]}{{sell_tot}}+{YC[y]}{{adm_tot}}+{YC[y]}{{rd_tot}}", fcst=lambda c: f"={L('sell_tot',c)}+{L('adm_tot',c)}+{L('rd_tot',c)}", bold=True)
op.row('opex_r', '  期间费用率', '%', hist=lambda y: f"=IFERROR({YC[y]}{{opex_tot}}/{RAW(HL,'营业收入',y)},0)", fcst=lambda c: f"=IFERROR({L('opex_tot',c)}/{G(S_REV,'rev_tot',c)},0)", fmt=PCT)
op.row('staff_r', '  职工薪酬总额/收入', '%', hist=lambda y: (f"=IFERROR({YC[y]}{{comp_tot}}/{RAW(HL,'营业收入',y)},0)" if DH('comp_total') and y in DH('comp_total') else None), fcst=lambda c: f"=IFERROR({L('comp_tot',c)}/{G(S_REV,'rev_tot',c)},0)", fmt=PCT)

# ---------------------------------------------------------------- 资本开支与折旧
cp = Sheet(S_CAP, '资本开支、在建工程与固定资产折旧（按资产类别）',
           '固定资产变动表来自年报附注（合并）；预测：新增资产按类别结构分配，折旧=（期初原值+50%新增-50%处置）×类别年折旧率。')
cp.sec('A. 资本开支计划')
cp.row('cap_maint', '维护性资本开支（上年固定资产折旧×维护比例）', '百万元', fcst=lambda c: f"={L('fa_dep',prev(c))}*{A('maint_ratio',c)}")
cp.row('cap_dom', '国内扩产（MLCC高容、光通信、陶瓷材料等）', '百万元', fcst=lambda c: f"={A('cap_dom',c)}")
cp.row('cap_hpool', '  H股募资中境外新建扩建项目资金池', '百万元', fcst=lambda c: f"='{S_SH}'!$H${REG[(S_SH,'h_net')]}*{A1('h_os_pct')}" if c == 'H' else None,
       src='H股招股章程：所得款项净额约41.2%投向泰国、德国新建扩建项目（燃料电池扩建、高精度压电微点胶、数据中心电子元件、通信器件）')
cp.row('cap_h', 'H股境外项目（泰国、德国）', '百万元', fcst=lambda c: f"=$H{REG[(S_CAP,'cap_hpool')]}*{A('h_phase',c)}")
cp.row('cap_int', '土地使用权、软件及长期待摊', '百万元', fcst=lambda c: f"={A('cap_int',c)}")
cp.row('capex_tot', '资本开支合计（历史=购建长期资产支付的现金；预测=当年新增长期资产投入，假设当期付现）', '百万元', hist=lambda y: f"={RAW(HCF,'购建固定资产、无形资产和其他长期资产支付的现金',y)}",
       fcst=lambda c: f"=({L('cap_maint',c)}+{L('cap_dom',c)}+{L('cap_h',c)}+{L('cap_int',c)})*(1+{A1('s_capex')}" + (f"*{A1('w26')}" if c == 'H' else '') + ")", bold=True)
cp.row('cap_cip', '  其中：投入在建工程（历史为权责口径的在建工程增加，与上行现金口径不勾稽）', '百万元', hist=lambda y: (f"={YC[y]}{{cip_add}}" if y in H23 else None), fcst=lambda c: f"=({L('capex_tot',c)}-{L('cap_int',c)})*{A('cip_share',c)}")
cp.row('cap_fa', '  其中：直接购置固定资产（历史为权责口径）', '百万元', hist=DH('fa_total', 'add_purchase'), fcst=lambda c: f"={L('capex_tot',c)}-{L('cap_int',c)}-{L('cap_cip',c)}")
cp.row('capex_r', '  资本开支/收入', '%', hist=lambda y: f"=IFERROR({YC[y]}{{capex_tot}}/{RAW(HL,'营业收入',y)},0)", fcst=lambda c: f"=IFERROR({L('capex_tot',c)}/{G(S_REV,'rev_tot',c)},0)", fmt=PCT)
cp.row('capex_da', '  资本开支/折旧摊销', 'x', hist=lambda y: f"=IFERROR({YC[y]}{{capex_tot}}/{YC[y]}{{da_tot}},0)", fcst=lambda c: f"=IFERROR({L('capex_tot',c)}/{L('da_tot',c)},0)", fmt='0.00"x"')
cp.sec('B. 在建工程')
cp.row('cip_open', '期初余额', '百万元', hist=lambda y: (f"={RAW(HB,'在建工程(合计)',{'2022A':'2021A','2023A':'2022A','2024A':'2023A','2025A':'2024A'}[y])}" if y != '2021A' else None), fcst=lambda c: f"={L('cip_close',prev(c))}")
cp.row('cip_add', '本期增加（资本开支投入）', '百万元', hist=lambda y: (f"={YC[y]}{{cip_close}}-{YC[y]}{{cip_open}}+{YC[y]}{{cip_tr}}" if y in H23 else None), fcst=lambda c: f"={L('cap_cip',c)}")
cp.row('cip_tr', '转入固定资产', '百万元', hist=DH('fa_total', 'add_cip'), fcst=lambda c: f"={L('cip_open',c)}*{A('cip_t_open',c)}+{L('cip_add',c)}*{A('cip_t_new',c)}")
cp.row('cip_close', '期末余额', '百万元', hist=lambda y: f"={RAW(HB,'在建工程(合计)',y)}", fcst=lambda c: f"={L('cip_open',c)}+{L('cip_add',c)}-{L('cip_tr',c)}", bold=True)
cp.row('cip_turn', '  转固率（转入/（期初+增加））', '%', hist=lambda y: (f"=IFERROR({YC[y]}{{cip_tr}}/({YC[y]}{{cip_open}}+{YC[y]}{{cip_add}}),0)" if y in H23 else None), fcst=lambda c: f"=IFERROR({L('cip_tr',c)}/({L('cip_open',c)}+{L('cip_add',c)}),0)", fmt=PCT)
CATS = [('bld', '房屋及建筑物'), ('mach', '机器设备'), ('trans', '运输设备'), ('elec', '电子设备及其他')]
cp.sec('C. 固定资产（按类别，原值/累计折旧）')
for k, nm in CATS:
    cp.hdr([f'{nm}', ''] + YEARS + [''])
    cp.row(f'{k}_go', f'  原值期初', '百万元', hist=lambda y, nm=nm: (D('fa', nm, 'gross_open', FY[y]) if D('fa', nm, 'gross_open', FY[y]) is not None else None), fcst=lambda c, k=k: f"={L(k+'_gc',prev(c))}")
    cp.row(f'{k}_add', f'  本期新增（在建转入+购置+其他）', '百万元', hist=lambda y, nm=nm: (round(sum(v for v in [D('fa', nm, 'add_cip', FY[y]), D('fa', nm, 'add_purchase', FY[y]), D('fa', nm, 'add_other', FY[y])] if v), 6) if D('fa', nm, 'gross_open', FY[y]) is not None else None),
           fcst=lambda c, k=k: f"=({L('cip_tr',c)}+{L('cap_fa',c)})*{A('mix_'+k,c)}")
    cp.row(f'{k}_disp', f'  本期减少（处置报废及其他）', '百万元', hist=lambda y, nm=nm: (round(sum(v for v in [D('fa', nm, 'disp_gross', FY[y]), D('fa', nm, 'red_other', FY[y])] if v), 6) if D('fa', nm, 'gross_open', FY[y]) is not None else None),
           fcst=lambda c, k=k: f"={L(k+'_go',c)}*{A('disp_'+k,c)}")
    cp.row(f'{k}_gc', f'  原值期末', '百万元', hist=lambda y, nm=nm: D('fa', nm, 'gross_close', FY[y]), fcst=lambda c, k=k: f"={L(k+'_go',c)}+{L(k+'_add',c)}-{L(k+'_disp',c)}", bold=True)
    cp.row(f'{k}_dep', f'  本期折旧', '百万元', hist=lambda y, nm=nm: D('fa', nm, 'dep', FY[y]),
           fcst=lambda c, k=k: f"=({L(k+'_go',c)}+0.5*{L(k+'_add',c)}-0.5*{L(k+'_disp',c)})*{A('dep_'+k,c)}")
    cp.row(f'{k}_rate', f'  折旧率（折旧/平均原值）', '%', hist=lambda y, k=k, nm=nm: (f"=IFERROR({YC[y]}{{{k}_dep}}/({YC[y]}{{{k}_go}}+0.5*{YC[y]}{{{k}_add}}-0.5*{YC[y]}{{{k}_disp}}),0)" if D('fa', nm, 'dep', FY[y]) is not None else None),
           fcst=lambda c, k=k: f"={A('dep_'+k,c)}", fmt=PCT2)
    cp.row(f'{k}_ado', f'  累计折旧期初', '百万元', hist=lambda y, nm=nm: D('fa', nm, 'ad_open', FY[y]), fcst=lambda c, k=k: f"={L(k+'_adc',prev(c))}")
    cp.row(f'{k}_add_disp', f'  累计折旧减少（处置及其他）', '百万元', hist=lambda y, nm=nm: (round(sum(v for v in [D('fa', nm, 'ad_disp', FY[y]), D('fa', nm, 'ad_other_net', FY[y])] if v), 6) if D('fa', nm, 'ad_open', FY[y]) is not None else None),
           fcst=lambda c, k=k: f"={L(k+'_disp',c)}*{A('adr_'+k,c)}")
    cp.row(f'{k}_adc', f'  累计折旧期末', '百万元', hist=lambda y, nm=nm: D('fa', nm, 'ad_close', FY[y]), fcst=lambda c, k=k: f"={L(k+'_ado',c)}+{L(k+'_dep',c)}-{L(k+'_add_disp',c)}")
    cp.row(f'{k}_nbv', f'  账面净值（未扣减值）', '百万元', hist=lambda y, k=k, nm=nm: (f"={YC[y]}{{{k}_gc}}-{YC[y]}{{{k}_adc}}" if D('fa', nm, 'gross_close', FY[y]) is not None else None), fcst=lambda c, k=k: f"={L(k+'_gc',c)}-{L(k+'_adc',c)}")
    cp.row(f'{k}_age', f'  净值/原值（新旧程度）', '%', hist=lambda y, k=k, nm=nm: (f"=IFERROR({YC[y]}{{{k}_nbv}}/{YC[y]}{{{k}_gc}},0)" if D('fa', nm, 'gross_close', FY[y]) is not None else None), fcst=lambda c, k=k: f"=IFERROR({L(k+'_nbv',c)}/{L(k+'_gc',c)},0)", fmt=PCT)
cp.sec('D. 固定资产合计')
catsum = lambda suf, c: '+'.join(L(k + suf, c) for k, _ in CATS)
cp.row('fa_gross', '固定资产原值', '百万元', hist=lambda y: (f"=" + '+'.join(f"{YC[y]}{{{k}_gc}}" for k, _ in CATS) if D('fa', '机器设备', 'gross_close', FY[y]) is not None else None), fcst=lambda c: '=' + catsum('_gc', c))
cp.row('fa_ad', '累计折旧', '百万元', hist=lambda y: (f"=" + '+'.join(f"{YC[y]}{{{k}_adc}}" for k, _ in CATS) if D('fa', '机器设备', 'ad_close', FY[y]) is not None else None), fcst=lambda c: '=' + catsum('_adc', c))
cp.row('fa_imp', '减值准备', '百万元', hist=DH('fa_total', 'imp_close'), fcst=lambda c: f"={L('fa_imp',prev(c))}-{G(S_TAX,'imp_fa',c)}", src='加上当年固定资产减值损失（见税务与其他损益）')
cp.row('fa_clr', '固定资产清理', '百万元', hist=DH('fa_clear'), src='仅2023年末有余额1.92')
cp.row('fa_nbv', '固定资产账面价值（含清理）', '百万元', hist=lambda y: f"={RAW(HB,'固定资产(合计)',y)}", fcst=lambda c: f"={L('fa_gross',c)}-{L('fa_ad',c)}-{L('fa_imp',c)}", bold=True)
cp.row('fa_chk', '  校验：类别合计-资产负债表（历史）', '百万元', hist=lambda y: (f"=ROUND({YC[y]}{{fa_gross}}-{YC[y]}{{fa_ad}}-{YC[y]}{{fa_imp}}+N({YC[y]}{{fa_clr}})-{YC[y]}{{fa_nbv}},1)" if D('fa', '机器设备', 'gross_close', FY[y]) is not None else None))
cp.row('fa_dep', '固定资产折旧合计', '百万元', hist=lambda y: (f"=" + '+'.join(f"{YC[y]}{{{k}_dep}}" for k, _ in CATS) if D('fa', '机器设备', 'dep', FY[y]) is not None else f"={RAW(HCF,'固定资产折旧、油气资产折耗、生产性生物资产折旧',y)}"),
       fcst=lambda c: '=' + catsum('_dep', c), bold=True)
cp.row('disp_nbv', '处置资产账面净值（假设按账面价值收回现金）', '百万元', fcst=lambda c: '=' + '+'.join(f"({L(k+'_disp',c)}-{L(k+'_add_disp',c)})" for k, _ in CATS))
cp.row('ppe', '固定资产+在建工程', '百万元', hist=lambda y: f"={YC[y]}{{fa_nbv}}+{YC[y]}{{cip_close}}", fcst=lambda c: f"={L('fa_nbv',c)}+{L('cip_close',c)}", bold=True)
cp.row('fa_turn', '  固定资产周转率（收入/平均原值）', 'x', hist=lambda y: (f"=IFERROR({RAW(HL,'营业收入',y)}/AVERAGE({YC[y]}{{fa_gross}},{prev(YC[y])}{{fa_gross}}),0)" if D('fa', '机器设备', 'gross_close', FY[y]) is not None and D('fa', '机器设备', 'gross_close', FY[{'2022A':'2021A','2023A':'2022A','2024A':'2023A','2025A':'2024A'}.get(y,'2021A')]) is not None else None),
       fcst=lambda c: f"=IFERROR({G(S_REV,'rev_tot',c)}/AVERAGE({L('fa_gross',c)},{L('fa_gross',prev(c))}),0)", fmt='0.00"x"')
cp.sec('E. 无形资产、使用权资产及其他长期资产')
cp.row('int_open', '无形资产及长期待摊期初', '百万元', hist=lambda y: (f"={prev(YC[y])}{{int_close}}" if y != '2021A' else None), fcst=lambda c: f"={L('int_close',prev(c))}")
cp.row('int_add', '  本期增加（资本开支）', '百万元', fcst=lambda c: f"={L('cap_int',c)}")
cp.row('int_amort', '  本期摊销', '百万元', hist=lambda y: f"={RAW(HCF,'无形资产摊销',y)}+{RAW(HCF,'长期待摊费用摊销',y)}", fcst=lambda c: f"={L('int_open',c)}*{A('int_amort_r',c)}")
cp.row('int_close', '无形资产及长期待摊期末', '百万元', hist=lambda y: f"={RAW(HB,'无形资产',y)}+N({RAW(HB,'长期待摊费用',y)})", fcst=lambda c: f"={L('int_open',c)}+{L('int_add',c)}-{L('int_amort',c)}", bold=True)
cp.row('rou', '使用权资产（假设续租维持不变）', '百万元', hist=lambda y: f"=N({RAW(HB,'使用权资产',y)})", fcst=lambda c: f"={L('rou',prev(c))}")
cp.row('rou_dep', '  使用权资产折旧', '百万元', hist=lambda y: f"=N({RAW(HCF,'使用权资产折旧',y)})", fcst=lambda c: f"={A('rou_dep',c)}")
cp.row('olta_flat', '商誉、投资性房地产及股权投资（维持不变）', '百万元', hist=lambda y: RAWS(HB, ['长期股权投资', '其他权益工具投资', '投资性房地产', '商誉'], y), fcst=lambda c: f"={L('olta_flat',prev(c))}")
cp.sec('F. 折旧摊销汇总与分摊')
cp.row('da_tot', '折旧摊销合计', '百万元', hist=lambda y: f"={YC[y]}{{fa_dep}}+{YC[y]}{{int_amort}}+{YC[y]}{{rou_dep}}", fcst=lambda c: f"={L('fa_dep',c)}+{L('int_amort',c)}+{L('rou_dep',c)}", bold=True)
cp.row('da_r', '  折旧摊销/收入', '%', hist=lambda y: f"=IFERROR({YC[y]}{{da_tot}}/{RAW(HL,'营业收入',y)},0)", fcst=lambda c: f"=IFERROR({L('da_tot',c)}/{G(S_REV,'rev_tot',c)},0)", fmt=PCT)
cp.row('da_cogs', '  计入营业成本', '百万元', fcst=lambda c: f"={L('da_tot',c)}*{A1('da_cogs')}")
cp.row('da_rd', '  计入研发费用', '百万元', fcst=lambda c: f"={L('da_tot',c)}*{A1('da_rd')}")
cp.row('da_sell', '  计入销售费用', '百万元', fcst=lambda c: f"={L('da_tot',c)}*{A1('da_sell')}")
cp.row('da_adm', '  计入管理费用（倒挤）', '百万元', fcst=lambda c: f"={L('da_tot',c)}-{L('da_cogs',c)}-{L('da_rd',c)}-{L('da_sell',c)}")

# ---------------------------------------------------------------- 营运资金
wcs = Sheet(S_WC, '营运资金明细（周转天数驱动）', '历史取Wind资产负债表；存货分类来自年报附注。周转天数按期末余额计算。')
wcs.sec('A. 经营性流动资产')
def wrow(key, label, rawlabs, fcst, fmt=NUM, src=None, bold=False):
    wcs.row(key, label, '百万元', hist=(lambda y: RAWS(HB, rawlabs, y)) if rawlabs else None, fcst=fcst, fmt=fmt, src=src, bold=bold)
REV = lambda c: G(S_REV, 'rev_tot', c)
COGS = lambda c: G(S_COST, 'cogs_tot', c)
wrow('nrec', '应收票据', ['应收票据'], lambda c: f"={REV(c)}*{A('d_nrec',c)}/365")
wrow('ar', '应收账款', ['应收账款'], lambda c: f"={REV(c)}*({A('d_ar',c)}+{A1('s_wcdays')})/365")
wrow('rfin', '应收款项融资', ['应收款项融资'], lambda c: f"={REV(c)}*{A('d_rfin',c)}/365")
wrow('pre', '预付款项', ['预付款项'], lambda c: f"={COGS(c)}*{A('r_pre',c)}")
wrow('orec', '其他应收款', ['其他应收款(合计)'], lambda c: f"={REV(c)}*{A('r_orec',c)}")
wcs.row('inv_raw', '  存货-原材料（账面余额）', '百万元', hist=DH('inv', '原材料'), fcst=lambda c: f"={G(S_COST,'mat',c)}*{A('d_inv_raw',c)}/365")
wcs.row('inv_wip', '  存货-在产品及半成品（账面余额）', '百万元', hist=DH('inv', '在产品'), fcst=lambda c: f"={COGS(c)}*{A('d_inv_wip',c)}/365")
wcs.row('inv_fg', '  存货-库存商品及发出商品（账面余额）', '百万元', hist=DH('inv', '库存商品及发出商品'), fcst=lambda c: f"={COGS(c)}*({A('d_inv_fg',c)}+{A1('s_wcdays')})/365")
wcs.row('inv_oth', '  存货-其他（周转材料、委托加工等）', '百万元', hist=DH('inv', '其他'), fcst=lambda c: f"={COGS(c)}*{A('r_inv_oth',c)}")
wcs.row('inv_prov', '  存货跌价准备', '百万元', hist=DH('inv', '跌价准备'), fcst=lambda c: f"=({L('inv_raw',c)}+{L('inv_wip',c)}+{L('inv_fg',c)}+{L('inv_oth',c)})*{A('r_inv_prov',c)}")
wrow('inv', '存货（账面价值）', ['存货'], lambda c: f"={L('inv_raw',c)}+{L('inv_wip',c)}+{L('inv_fg',c)}+{L('inv_oth',c)}-{L('inv_prov',c)}", bold=True)
wcs.row('inv_chk', '  校验：分类合计-资产负债表（历史）', '百万元', hist=lambda y: (f"=ROUND({YC[y]}{{inv_raw}}+{YC[y]}{{inv_wip}}+{YC[y]}{{inv_fg}}+{YC[y]}{{inv_oth}}-{YC[y]}{{inv_prov}}-{YC[y]}{{inv}},1)" if DH('inv', '原材料') and y in DH('inv', '原材料') else None))
wrow('oca', '其他流动资产', ['其他流动资产'], lambda c: f"={REV(c)}*{A('r_oca',c)}")
wcs.row('op_ca', '经营性流动资产合计', '百万元', hist=lambda y: f"={YC[y]}{{nrec}}+{YC[y]}{{ar}}+{YC[y]}{{rfin}}+{YC[y]}{{pre}}+{YC[y]}{{orec}}+{YC[y]}{{inv}}+{YC[y]}{{oca}}",
        fcst=lambda c: f"={L('nrec',c)}+{L('ar',c)}+{L('rfin',c)}+{L('pre',c)}+{L('orec',c)}+{L('inv',c)}+{L('oca',c)}", bold=True)
wcs.sec('B. 经营性流动负债')
wrow('npay', '应付票据', ['应付票据'], lambda c: f"={COGS(c)}*{A('d_npay',c)}/365")
wrow('ap', '应付账款', ['应付账款'], lambda c: f"={COGS(c)}*{A('d_ap',c)}/365")
wrow('contract', '合同负债及预收款项', ['预收款项', '合同负债'], lambda c: f"={REV(c)}*{A('r_contract',c)}")
wrow('payroll', '应付职工薪酬', ['应付职工薪酬'], lambda c: f"={G(S_OPEX,'comp_tot',c)}*{A('r_payroll',c)}")
wrow('taxpay', '应交税费', ['应交税费'], lambda c: f"={REV(c)}*{A('r_taxpay',c)}")
wrow('othpay', '其他应付款', ['其他应付款(合计)'], lambda c: f"={REV(c)}*{A('r_othpay',c)}")
wrow('ocl', '其他流动负债', ['其他流动负债'], lambda c: f"={REV(c)}*{A('r_ocl',c)}")
wcs.row('op_cl', '经营性流动负债合计', '百万元', hist=lambda y: f"={YC[y]}{{npay}}+{YC[y]}{{ap}}+{YC[y]}{{contract}}+{YC[y]}{{payroll}}+{YC[y]}{{taxpay}}+{YC[y]}{{othpay}}+{YC[y]}{{ocl}}",
        fcst=lambda c: f"={L('npay',c)}+{L('ap',c)}+{L('contract',c)}+{L('payroll',c)}+{L('taxpay',c)}+{L('othpay',c)}+{L('ocl',c)}", bold=True)
wcs.sec('C. 营运资本与周转')
wcs.row('nwc', '营运资本（经营性流动资产-经营性流动负债）', '百万元', hist=lambda y: f"={YC[y]}{{op_ca}}-{YC[y]}{{op_cl}}", fcst=lambda c: f"={L('op_ca',c)}-{L('op_cl',c)}", bold=True)
wcs.row('dnwc', '营运资本增加', '百万元', hist=lambda y: (f"={YC[y]}{{nwc}}-{prev(YC[y])}{{nwc}}" if y != '2021A' else None), fcst=lambda c: f"={L('nwc',c)}-{L('nwc',prev(c))}")
wcs.row('nwc_r', '  营运资本/收入', '%', hist=lambda y: f"=IFERROR({YC[y]}{{nwc}}/{RAW(HL,'营业收入',y)},0)", fcst=lambda c: f"=IFERROR({L('nwc',c)}/{REV(c)},0)", fmt=PCT)
wcs.row('dso', '应收周转天数（应收票据+账款+融资）', '天', hist=lambda y: f"=IFERROR(({YC[y]}{{nrec}}+{YC[y]}{{ar}}+{YC[y]}{{rfin}})/{RAW(HL,'营业收入',y)}*365,0)", fcst=lambda c: f"=IFERROR(({L('nrec',c)}+{L('ar',c)}+{L('rfin',c)})/{REV(c)}*365,0)", fmt=DAYS)
wcs.row('dio', '存货周转天数', '天', hist=lambda y: f"=IFERROR({YC[y]}{{inv}}/{RAW(HL,'营业成本',y)}*365,0)", fcst=lambda c: f"=IFERROR({L('inv',c)}/{COGS(c)}*365,0)", fmt=DAYS)
wcs.row('dpo', '应付周转天数（应付票据+账款）', '天', hist=lambda y: f"=IFERROR(({YC[y]}{{npay}}+{YC[y]}{{ap}})/{RAW(HL,'营业成本',y)}*365,0)", fcst=lambda c: f"=IFERROR(({L('npay',c)}+{L('ap',c)})/{COGS(c)}*365,0)", fmt=DAYS)
wcs.row('ccc', '现金转换周期', '天', hist=lambda y: f"={YC[y]}{{dso}}+{YC[y]}{{dio}}-{YC[y]}{{dpo}}", fcst=lambda c: f"={L('dso',c)}+{L('dio',c)}-{L('dpo',c)}", fmt=DAYS, bold=True)
wcs.row('d_ar_h', '  应收账款天数', '天', hist=lambda y: f"=IFERROR({YC[y]}{{ar}}/{RAW(HL,'营业收入',y)}*365,0)", fcst=lambda c: f"=IFERROR({L('ar',c)}/{REV(c)}*365,0)", fmt=DAYS)
wcs.row('d_nrec_h', '  应收票据天数', '天', hist=lambda y: f"=IFERROR({YC[y]}{{nrec}}/{RAW(HL,'营业收入',y)}*365,0)", fcst=lambda c: f"=IFERROR({L('nrec',c)}/{REV(c)}*365,0)", fmt=DAYS)
wcs.row('d_rfin_h', '  应收款项融资天数', '天', hist=lambda y: f"=IFERROR({YC[y]}{{rfin}}/{RAW(HL,'营业收入',y)}*365,0)", fcst=lambda c: f"=IFERROR({L('rfin',c)}/{REV(c)}*365,0)", fmt=DAYS)
wcs.row('d_raw_h', '  原材料天数（按原材料成本）', '天', hist=lambda y: (f"=IFERROR({YC[y]}{{inv_raw}}/{G(S_COST,'mat',YC[y])}*365,0)" if DH('inv','原材料') and y in DH('inv','原材料') and DH('cogs_nature','原材料') and y in DH('cogs_nature','原材料') else None), fcst=lambda c: f"=IFERROR({L('inv_raw',c)}/{G(S_COST,'mat',c)}*365,0)", fmt=DAYS)
wcs.row('d_wip_h', '  在产品及半成品天数', '天', hist=lambda y: (f"=IFERROR({YC[y]}{{inv_wip}}/{RAW(HL,'营业成本',y)}*365,0)" if DH('inv','在产品') and y in DH('inv','在产品') else None), fcst=lambda c: f"=IFERROR({L('inv_wip',c)}/{COGS(c)}*365,0)", fmt=DAYS)
wcs.row('d_fg_h', '  库存商品及发出商品天数', '天', hist=lambda y: (f"=IFERROR({YC[y]}{{inv_fg}}/{RAW(HL,'营业成本',y)}*365,0)" if DH('inv','库存商品及发出商品') and y in DH('inv','库存商品及发出商品') else None), fcst=lambda c: f"=IFERROR({L('inv_fg',c)}/{COGS(c)}*365,0)", fmt=DAYS)
wcs.row('d_npay_h', '  应付票据天数', '天', hist=lambda y: f"=IFERROR({YC[y]}{{npay}}/{RAW(HL,'营业成本',y)}*365,0)", fcst=lambda c: f"=IFERROR({L('npay',c)}/{COGS(c)}*365,0)", fmt=DAYS)
wcs.row('d_ap_h', '  应付账款天数', '天', hist=lambda y: f"=IFERROR({YC[y]}{{ap}}/{RAW(HL,'营业成本',y)}*365,0)", fcst=lambda c: f"=IFERROR({L('ap',c)}/{COGS(c)}*365,0)", fmt=DAYS)
wcs.row('prov_r_h', '  存货跌价准备/存货余额', '%', hist=lambda y: (f"=IFERROR({YC[y]}{{inv_prov}}/({YC[y]}{{inv_raw}}+{YC[y]}{{inv_wip}}+{YC[y]}{{inv_fg}}+{YC[y]}{{inv_oth}}),0)" if DH('inv','跌价准备') and y in DH('inv','跌价准备') else None), fcst=lambda c: f"={A('r_inv_prov',c)}", fmt=PCT)

# ---------------------------------------------------------------- 融资与资金
fn = Sheet(S_FIN, '融资、资金与权益（利息按期初/平均余额，无循环引用）', '货币资金由“三表预测”现金流倒挤；利息收入按期初余额计息以避免循环；有息负债按假设的净借款变动滚动。')
fn.sec('A. 货币资金、理财与定期存款')
fn.row('cash', '货币资金', '百万元', hist=lambda y: f"={RAW(HB,'货币资金',y)}", fcst=lambda c: f"={G(S_M,'cash',c)}", bold=True)
fn.row('rcash', '  其中：受限资金', '百万元', hist=DH('cash', '受限资金'), fcst=lambda c: f"={L('rcash',prev(c))}", src='年报货币资金附注：使用受限的货币资金（票据保证金等）')
fn.row('ucash', '  其中：可自由使用', '百万元', hist=lambda y: (f"={YC[y]}{{cash}}-{YC[y]}{{rcash}}" if DH('cash','受限资金') and y in DH('cash','受限资金') else None), fcst=lambda c: f"={L('cash',c)}-{L('rcash',c)}")
fn.row('tfa', '交易性金融资产（银行理财、结构性存款）', '百万元', hist=lambda y: f"={RAW(HB,'交易性金融资产',y)}", fcst=lambda c: f"={L('tfa',prev(c))}+{A('tfa_net',c)}")
fn.row('td', '其他非流动资产-定期存款/大额存单', '百万元', hist=DH('onca', '定期存款及大额存单'), fcst=lambda c: f"={L('td',prev(c))}+{A('td_net',c)}")
fn.row('onca_oth', '其他非流动资产-预付工程设备款及其他', '百万元', hist=lambda y: (f"={RAW(HB,'其他非流动资产',y)}-{YC[y]}{{td}}" if DH('onca','定期存款及大额存单') and y in DH('onca','定期存款及大额存单') else None), fcst=lambda c: f"={L('onca_oth',prev(c))}")
fn.row('onca', '其他非流动资产合计', '百万元', hist=lambda y: f"={RAW(HB,'其他非流动资产',y)}", fcst=lambda c: f"={L('td',c)}+{L('onca_oth',c)}", bold=True)
fn.sec('B. 利息收入与理财收益')
fn.row('int_cash', '活期及短期存款利息', '百万元', fcst=lambda c: f"={L('ucash',prev(c))}*{A('y_cash',c)}")
fn.row('int_td', '定期存款/大额存单利息', '百万元', fcst=lambda c: f"={L('td',prev(c))}*{A('y_td',c)}")
fn.row('int_h', 'H股募资款当年利息（按在外天数）', '百万元', fcst=lambda c: f"='{S_SH}'!{c}{REG[(S_SH,'h_net')]}*{A('y_cash',c)}*{A1('h_days')}/365")
fn.row('int_inc', '利息收入合计', '百万元', hist=lambda y: f"={RAW(HL,'减：利息收入',y)}", fcst=lambda c: f"={L('int_cash',c)}+{L('int_td',c)}+{L('int_h',c)}", bold=True)
fn.row('yld_hist', '  综合利息收益率（利息收入/期初货币资金+其他非流动资产）', '%', hist=lambda y: (f"=IFERROR({YC[y]}{{int_inc}}/({prev(YC[y])}{{cash}}+{prev(YC[y])}{{onca}}),0)" if y != '2021A' else None),
       fcst=lambda c: f"=IFERROR({L('int_inc',c)}/({L('cash',prev(c))}+{L('onca',prev(c))}),0)", fmt=PCT2)
fn.row('inv_inc', '投资收益+公允价值变动收益', '百万元', hist=lambda y: RAWS(HL, ['投资净收益', '公允价值变动净收益'], y), fcst=lambda c: f"={L('tfa',prev(c))}*{A('y_tfa',c)}", bold=True)
fn.sec('C. 有息负债与利息支出')
fn.row('st_debt', '短期借款', '百万元', hist=lambda y: f"=N({RAW(HB,'短期借款',y)})", fcst=lambda c: f"={L('st_debt',prev(c))}+{A('st_chg',c)}")
fn.row('lt_debt', '长期借款（含一年内到期）', '百万元', hist=lambda y: RAWS(HB, ['长期借款', '一年内到期的非流动负债'], y), fcst=lambda c: f"={L('lt_debt',prev(c))}+{A('lt_chg',c)}")
fn.row('debt', '有息负债合计', '百万元', hist=lambda y: f"={YC[y]}{{st_debt}}+{YC[y]}{{lt_debt}}", fcst=lambda c: f"={L('st_debt',c)}+{L('lt_debt',c)}", bold=True)
fn.row('lease', '租赁负债（假设维持不变）', '百万元', hist=lambda y: f"=N({RAW(HB,'租赁负债',y)})", fcst=lambda c: f"={L('lease',prev(c))}")
fn.row('int_exp', '利息费用', '百万元', hist=lambda y: f"={RAW(HL,'其中：利息费用',y)}",
       fcst=lambda c: f"=AVERAGE({L('st_debt',prev(c))},{L('st_debt',c)})*{A('r_st',c)}+AVERAGE({L('lt_debt',prev(c))},{L('lt_debt',c)})*{A('r_lt',c)}+{L('lease',c)}*{A('r_lease',c)}", bold=True)
fn.row('rate_hist', '  平均借款成本（利息费用/平均有息负债+租赁负债）', '%', hist=lambda y: (f"=IFERROR({YC[y]}{{int_exp}}/AVERAGE({prev(YC[y])}{{debt}}+{prev(YC[y])}{{lease}},{YC[y]}{{debt}}+{YC[y]}{{lease}}),0)" if y != '2021A' else None),
       fcst=lambda c: f"=IFERROR({L('int_exp',c)}/AVERAGE({L('debt',prev(c))}+{L('lease',prev(c))},{L('debt',c)}+{L('lease',c)}),0)", fmt=PCT2)
fn.row('fx_fee', '汇兑损益、手续费及其他', '百万元', hist=lambda y: f"={RAW(HL,'财务费用',y)}-{YC[y]}{{int_exp}}+{YC[y]}{{int_inc}}", fcst=lambda c: f"={A('fx_fee',c)}",
       src='历史倒挤：财务费用-利息费用+利息收入（主要为汇兑损益）；2026H1受汇率影响汇兑损失增加')
fn.row('fin_exp', '财务费用（负数为净收益）', '百万元', hist=lambda y: f"={RAW(HL,'财务费用',y)}", fcst=lambda c: f"={L('int_exp',c)}-{L('int_inc',c)}+{L('fx_fee',c)}", bold=True)
fn.sec('D. 所有者权益变动')
fn.row('sc', '股本', '百万元', hist=lambda y: f"={RAW(HB,'实收资本(或股本)',y)}", fcst=lambda c: f"={L('sc',prev(c))}+'{S_SH}'!{c}{REG[(S_SH,'h_new')]}")
fn.row('apic', '资本公积', '百万元', hist=lambda y: f"={RAW(HB,'资本公积',y)}", fcst=lambda c: f"={L('apic',prev(c))}+'{S_SH}'!{c}{REG[(S_SH,'h_net')]}-'{S_SH}'!{c}{REG[(S_SH,'h_new')]}")
fn.row('tsy', '减：库存股', '百万元', hist=lambda y: f"=N({RAW(HB,'减：库存股',y)})", fcst=lambda c: f"={L('tsy',prev(c))}+'{S_SH}'!{c}{REG[(S_SH,'bb_amt')]}")
fn.row('oci', '其他综合收益（假设不变）', '百万元', hist=lambda y: f"={RAW(HB,'其他综合收益',y)}", fcst=lambda c: f"={L('oci',prev(c))}")
fn.row('surplus', '盈余公积', '百万元', hist=lambda y: f"={RAW(HB,'盈余公积',y)}", fcst=lambda c: f"={L('surplus',prev(c))}",
       src='2025年末盈余公积1,131.1百万元已超过股本50%（958.2），母公司可不再提取法定盈余公积；故预测维持不变')
fn.row('re', '未分配利润', '百万元', hist=lambda y: f"={RAW(HB,'未分配利润',y)}", fcst=lambda c: f"={L('re',prev(c))}+{G(S_M,'np',c)}-'{S_SH}'!{c}{REG[(S_SH,'div')]}")
fn.row('eq_p', '归母所有者权益', '百万元', hist=lambda y: f"={RAW(HB,'归属于母公司所有者权益合计',y)}", fcst=lambda c: f"={L('sc',c)}+{L('apic',c)}-{L('tsy',c)}+{L('oci',c)}+{L('surplus',c)}+{L('re',c)}", bold=True)
fn.row('eq_chk', '  校验：权益明细合计-报告值（历史）', '百万元', hist=lambda y: f"=ROUND({YC[y]}{{sc}}+{YC[y]}{{apic}}-{YC[y]}{{tsy}}+{YC[y]}{{oci}}+{YC[y]}{{surplus}}+{YC[y]}{{re}}-{YC[y]}{{eq_p}},1)")
fn.sec('E. 分红、回购与杠杆')
fn.row('div', '当年派发现金股利', '百万元', hist=lambda y: (f"='{S_SH}'!{YC[y]}{REG[(S_SH,'div')]}" if y == '2025A' else None), fcst=lambda c: f"='{S_SH}'!{c}{REG[(S_SH,'div')]}")
fn.row('dps', '每股股利（按派息时流通股）', '元/股', fcst=lambda c: f"=IFERROR({L('div',c)}/'{S_SH}'!{prev(c)}{REG[(S_SH,'out')]},0)", fmt=PS)
fn.row('div_yield', '  股息率（当前股价）', '%', fcst=lambda c: f"=IFERROR({L('dps',c)}/{A1('price')},0)", fmt=PCT2)
fn.row('bb', '股份回购', '百万元', fcst=lambda c: f"='{S_SH}'!{c}{REG[(S_SH,'bb_amt')]}")
fn.row('netcash', '净现金（货币资金+理财+一年以上定期存款-有息负债-租赁负债）', '百万元', hist=lambda y: f"={YC[y]}{{cash}}+{YC[y]}{{tfa}}+N({YC[y]}{{td}})-{YC[y]}{{debt}}-{YC[y]}{{lease}}",
       fcst=lambda c: f"={L('cash',c)}+{L('tfa',c)}+{L('td',c)}-{L('debt',c)}-{L('lease',c)}", bold=True,
       src='一年以上定期存款及大额存单列示于其他非流动资产（年报附注），视同现金；2021A未拆分，按0计')
fn.row('icr', '利息保障倍数（EBIT/利息费用）', 'x', hist=lambda y: f"=IFERROR({G(S_M,'ebit',YC[y])}/{YC[y]}{{int_exp}},0)", fcst=lambda c: f"=IFERROR({G(S_M,'ebit',c)}/{L('int_exp',c)},0)", fmt=MULT)
fn.row('de', '有息负债/归母权益', '%', hist=lambda y: f"=IFERROR({YC[y]}{{debt}}/{YC[y]}{{eq_p}},0)", fcst=lambda c: f"=IFERROR({L('debt',c)}/{L('eq_p',c)},0)", fmt=PCT)

# ---------------------------------------------------------------- 税务与其他损益
tx = Sheet(S_TAX, '税金、其他收益、减值与所得税桥', '所得税：法定15%（高新技术企业）为起点，叠加子公司税率差异、研发费用加计扣除与其他调整。')
tx.sec('A. 税金及附加')
tx.row('stax', '税金及附加', '百万元', hist=lambda y: f"={RAW(HL,'税金及附加',y)}", fcst=lambda c: f"={REV(c)}*{A('r_stax',c)}", bold=True)
tx.row('stax_r', '  占收入', '%', hist=lambda y: f"=IFERROR({YC[y]}{{stax}}/{RAW(HL,'营业收入',y)},0)", fcst=lambda c: f"=IFERROR({L('stax',c)}/{REV(c)},0)", fmt=PCT2)
tx.sec('B. 其他收益')
tx.row('oi_def', '与资产相关的政府补助摊销（递延收益转入）', '百万元', hist=DH('oi', '与资产相关的政府补助'), fcst=lambda c: f"={A('oi_def',c)}")
tx.row('oi_vat', '增值税加计抵减', '百万元', hist=DH('oi', '增值税加计抵减'), fcst=lambda c: f"={G(S_COST,'mat',c)}*{A('vat_addon_r',c)}",
       src='先进制造业企业按当期可抵扣进项税额加计5%抵减（政策期至2027年底），按原材料成本比例近似')
tx.row('oi_grant', '与收益相关的政府补助及其他', '百万元', hist=lambda y: (f"={YC[y]}{{oi_tot}}-N({YC[y]}{{oi_def}})-N({YC[y]}{{oi_vat}})" if DH('oi','增值税加计抵减') and y in DH('oi','增值税加计抵减') else None), fcst=lambda c: f"={REV(c)}*{A('grant_r',c)}")
tx.row('oi_tot', '其他收益合计', '百万元', hist=lambda y: f"={RAW(HL,'加：其他收益',y)}", fcst=lambda c: f"={L('oi_def',c)}+{L('oi_vat',c)}+{L('oi_grant',c)}", bold=True)
tx.row('oi_r', '  占收入', '%', hist=lambda y: f"=IFERROR({YC[y]}{{oi_tot}}/{RAW(HL,'营业收入',y)},0)", fcst=lambda c: f"=IFERROR({L('oi_tot',c)}/{REV(c)},0)", fmt=PCT2)
tx.sec('C. 减值损失（负数为损失）')
tx.row('imp_inv', '资产减值损失-存货跌价（历史=资产减值损失合计）', '百万元', hist=lambda y: f"=N({RAW(HL,'资产减值损失',y)})",
       fcst=lambda c: f"=-({G(S_WC,'inv_raw',c)}+{G(S_WC,'inv_wip',c)}+{G(S_WC,'inv_fg',c)}+{G(S_WC,'inv_oth',c)})*{A('r_inv_wd',c)}")
tx.row('imp_cr', '信用减值损失', '百万元', hist=lambda y: f"=N({RAW(HL,'信用减值损失',y)})", fcst=lambda c: f"=-{G(S_WC,'ar',c)}*{A('r_credit',c)}")
tx.row('imp_fa', '固定资产减值损失', '百万元', fcst=lambda c: f"=-{A('imp_fa',c)}", src='2026H1计提19.44；历史各期无（历史列含于资产减值损失）')
tx.row('imp_tot', '减值损失合计', '百万元', hist=lambda y: f"={YC[y]}{{imp_inv}}+{YC[y]}{{imp_cr}}", fcst=lambda c: f"={L('imp_inv',c)}+{L('imp_cr',c)}+{L('imp_fa',c)}", bold=True)
tx.row('disp', '资产处置收益', '百万元', hist=lambda y: f"=N({RAW(HL,'资产处置收益',y)})", fcst=lambda c: f"=N({RAW(HL,'资产处置收益','2026H1')})" if c == 'H' else "=0",
       src='2026E取2026H1实际值（下半年假设为0）；2027E起为0')
tx.row('nonop', '营业外收支净额', '百万元', hist=lambda y: f"={RAW(HL,'加：营业外收入',y)}-{RAW(HL,'减：营业外支出',y)}", fcst=lambda c: f"={RAW(HL,'加：营业外收入','2026H1')}-{RAW(HL,'减：营业外支出','2026H1')}" if c == 'H' else "=0",
       src='同上')
tx.sec('D. 所得税桥')
tx.row('pbt', '利润总额', '百万元', hist=lambda y: f"={RAW(HL,'利润总额',y)}", fcst=lambda c: f"={G(S_M,'pbt',c)}", bold=True)
tx.row('t_stat', '按法定适用税率计算的所得税', '百万元', hist=lambda y: (f"={YC[y]}{{pbt}}*{A1('stat_rate')}" if y in H23 else None), fcst=lambda c: f"={L('pbt',c)}*{A1('stat_rate')}")
tx.row('t_diff', '子公司适用不同税率的影响', '百万元', hist=DH('tax', '子公司适用不同税率的影响'), fcst=lambda c: f"={L('pbt',c)}*{A('t_diff_r',c)}")
tx.row('t_rd', '研发费用加计扣除的影响', '百万元', hist=DH('tax', '研发费用加计扣除的影响'),
       fcst=lambda c: f"=-{G(S_OPEX,'rd_tot',c)}*{A('rd_elig',c)}*{A1('super_ratio')}*{A1('stat_rate')}")
tx.row('t_oth', '其他调整（不可抵扣、以前年度、非应税等）', '百万元',
       hist=lambda y: (f"={YC[y]}{{tax}}-{YC[y]}{{t_stat}}-N({YC[y]}{{t_diff}})-N({YC[y]}{{t_rd}})" if y in H23 else None), fcst=lambda c: f"={L('pbt',c)}*({A('t_oth_r',c)}+{A1('s_tax')}" + (f"*{A1('w26')}" if c == 'H' else '') + ")")
tx.row('tax', '所得税费用', '百万元', hist=lambda y: f"={RAW(HL,'减：所得税',y)}", fcst=lambda c: f"={L('t_stat',c)}+{L('t_diff',c)}+{L('t_rd',c)}+{L('t_oth',c)}", bold=True)
tx.row('etr', '  实际税率', '%', hist=lambda y: f"=IFERROR({YC[y]}{{tax}}/{YC[y]}{{pbt}},0)", fcst=lambda c: f"=IFERROR({L('tax',c)}/{L('pbt',c)},0)", fmt=PCT, bold=True)
for k, nm in (('t_diff', '税率差异'), ('t_rd', '加计扣除'), ('t_oth', '其他调整')):
    tx.row(k + '_pp', f'  {nm}对税率的影响', 'pp', hist=lambda y, k=k: (f"=IFERROR(N({YC[y]}{{{k}}})/{YC[y]}{{pbt}},0)" if y in H23 else None), fcst=lambda c, k=k: f"=IFERROR({L(k,c)}/{L('pbt',c)},0)", fmt=PCT2)

# ---------------------------------------------------------------- 股本与每股 (sheet built with engine)
sh = Sheet(S_SH, '股本、库存股、H股募资与分红', '期末流通股本=总股本-库存股，用于EPS、每股净资产与市场倍数；DCF每股价值按含库存股的总股本（全面摊薄）。H股于2026-07-09上市。')
sh.row('h_new', 'H股新发股数', '百万股', fcst=lambda c: f"={A1('h_sh')}/1e6" if c == 'H' else 0)
sh.row('h_gross', 'H股募资总额', '百万元', fcst=lambda c: f"={A1('h_sh')}*{A1('h_px')}*{A1('fx')}/1e6" if c == 'H' else 0)
sh.row('h_net', 'H股募资净额', '百万元', fcst=lambda c: f"={L('h_gross',c)}*(1-{A1('h_cost')})" if c == 'H' else 0, bold=True)
sh.row('total', '期末总股本', '百万股', hist={y: 1916.497371 for y in HIST}, fcst=lambda c: f"={L('total',prev(c))}+{L('h_new',c)}")
sh.row('bb_amt', '当年回购金额', '百万元', hist={'2025A': 175.440732}, fcst=lambda c: f"={A1('bb1_amt')}+{A1('bb2_amt')}" if c == 'H' else f"={A('bb_amt',c)}")
sh.row('bb_sh', '当年回购股数', '百万股', hist={'2025A': 5.1338}, fcst=lambda c: (f"={A1('bb1_sh')}/1e6+IFERROR({A1('bb2_amt')}/{A1('price')},0)" if c == 'H' else f"=IFERROR({L('bb_amt',c)}/{A1('price')},0)"))
sh.row('ts', '期末库存股', '百万股', hist={'2021A': 0, '2022A': 0, '2023A': 0, '2024A': 0, '2025A': 5.1338}, fcst=lambda c: f"={L('ts',prev(c))}+{L('bb_sh',c)}")
sh.row('out', '期末流通股本（扣除库存股）', '百万股', hist=lambda y: f"={YC[y]}{{total}}-{YC[y]}{{ts}}", fcst=lambda c: f"={L('total',c)}-{L('ts',c)}", bold=True)
sh.row('wavg', '加权平均股本（历史=归母净利润÷报告基本每股收益；预测H股按上市后天数加权）', '百万股', hist=lambda y: f"=IFERROR({G(S_M,'np',YC[y])}/{RAW(HL,'基本每股收益',y)},{YC[y]}{{out}})",
       fcst=lambda c: (f"={L('out',prev(c))}+{L('h_new',c)}*{A1('h_days')}/365-{A1('bb1_sh')}/1e6*{A1('bb1_months')}/12" if c == 'H' else f"=AVERAGE({L('out',prev(c))},{L('out',c)})"))
sh.row('div', '当年派发现金股利', '百万元', hist={'2025A': 726.318157}, fcst=lambda c: (f"={A1('div26')}" if c == 'H' else f"={A1('payout')}*{G(S_M,'np',prev(c))}"), bold=True)
sh.row('dps_d', '  对应每股股利', '元/股', hist={'2025A': 0.38}, fcst=lambda c: f"=IFERROR({L('div',c)}/{L('out',prev(c))},0)", fmt=PS)

# ---------------------------------------------------------------- 三表预测
m = Sheet(S_M, '三表联动预测（利润表、资产负债表、现金流量表）', '2021A-2025A链接历史报表；2026E-2030E链接各明细模块。货币资金由现金流量表倒挤，资产负债表自动配平。')
m.sec('利润表')
m.row('rev', '营业收入', hist=lambda y: f"={RAW(HL,'营业收入',y)}", fcst=lambda c: f"={G(S_REV,'rev_tot',c)}", bold=True)
m.row('rev_g', '  同比', '%', hist=lambda y: (f"=IFERROR({YC[y]}{{rev}}/{prev(YC[y])}{{rev}}-1,0)" if y != '2021A' else None), fcst=lambda c: f"=IFERROR({L('rev',c)}/{L('rev',prev(c))}-1,0)", fmt=PCT)
m.row('cogs', '营业成本', hist=lambda y: f"={RAW(HL,'营业成本',y)}", fcst=lambda c: f"={G(S_COST,'cogs_tot',c)}")
m.row('gp', '毛利', hist=lambda y: f"={YC[y]}{{rev}}-{YC[y]}{{cogs}}", fcst=lambda c: f"={L('rev',c)}-{L('cogs',c)}", bold=True)
m.row('gm', '  毛利率', '%', hist=lambda y: f"=IFERROR({YC[y]}{{gp}}/{YC[y]}{{rev}},0)", fcst=lambda c: f"=IFERROR({L('gp',c)}/{L('rev',c)},0)", fmt=PCT)
m.row('stax', '税金及附加', hist=lambda y: f"={RAW(HL,'税金及附加',y)}", fcst=lambda c: f"={G(S_TAX,'stax',c)}")
m.row('sell', '销售费用', hist=lambda y: f"={RAW(HL,'销售费用',y)}", fcst=lambda c: f"={G(S_OPEX,'sell_tot',c)}")
m.row('adm', '管理费用', hist=lambda y: f"={RAW(HL,'管理费用',y)}", fcst=lambda c: f"={G(S_OPEX,'adm_tot',c)}")
m.row('rd', '研发费用', hist=lambda y: f"={RAW(HL,'研发费用',y)}", fcst=lambda c: f"={G(S_OPEX,'rd_tot',c)}")
m.row('fin', '财务费用（负数为净收益）', hist=lambda y: f"={RAW(HL,'财务费用',y)}", fcst=lambda c: f"={G(S_FIN,'fin_exp',c)}")
m.row('oi', '其他收益', hist=lambda y: f"={RAW(HL,'加：其他收益',y)}", fcst=lambda c: f"={G(S_TAX,'oi_tot',c)}")
m.row('inv', '投资收益+公允价值变动', hist=lambda y: RAWS(HL, ['投资净收益', '公允价值变动净收益'], y), fcst=lambda c: f"={G(S_FIN,'inv_inc',c)}")
m.row('imp', '资产及信用减值损失', hist=lambda y: RAWS(HL, ['资产减值损失', '信用减值损失'], y), fcst=lambda c: f"={G(S_TAX,'imp_tot',c)}")
m.row('disp', '资产处置收益', hist=lambda y: f"=N({RAW(HL,'资产处置收益',y)})", fcst=lambda c: f"={G(S_TAX,'disp',c)}")
OPF = lambda c: f"={c}{{gp}}-{c}{{stax}}-{c}{{sell}}-{c}{{adm}}-{c}{{rd}}-{c}{{fin}}+{c}{{oi}}+{c}{{inv}}+{c}{{imp}}+{c}{{disp}}"
m.row('op', '营业利润', hist=lambda y: OPF(YC[y]), fcst=lambda c: OPF(c).replace('{', '<').replace('}', '>'), bold=True)
m.row('nonop', '营业外收支净额', hist=lambda y: f"={RAW(HL,'加：营业外收入',y)}-{RAW(HL,'减：营业外支出',y)}", fcst=lambda c: f"={G(S_TAX,'nonop',c)}")
m.row('pbt', '利润总额', hist=lambda y: f"={YC[y]}{{op}}+{YC[y]}{{nonop}}", fcst=lambda c: f"={L('op',c)}+{L('nonop',c)}", bold=True)
m.row('tax', '所得税', hist=lambda y: f"={RAW(HL,'减：所得税',y)}", fcst=lambda c: f"={G(S_TAX,'tax',c)}")
m.row('ni', '净利润', hist=lambda y: f"={YC[y]}{{pbt}}-{YC[y]}{{tax}}", fcst=lambda c: f"={L('pbt',c)}-{L('tax',c)}", bold=True)
m.row('mi', '少数股东损益', hist=lambda y: f"={RAW(HL,'减：少数股东损益',y)}", fcst=lambda c: f"={L('ni',c)}*{A1('minority')}")
m.row('np', '归母净利润', hist=lambda y: f"={YC[y]}{{ni}}-{YC[y]}{{mi}}", fcst=lambda c: f"={L('ni',c)}-{L('mi',c)}", bold=True)
m.row('np_g', '  同比', '%', hist=lambda y: (f"=IFERROR({YC[y]}{{np}}/{prev(YC[y])}{{np}}-1,0)" if y != '2021A' else None), fcst=lambda c: f"=IFERROR({L('np',c)}/{L('np',prev(c))}-1,0)", fmt=PCT)
m.row('npm', '  归母净利率', '%', hist=lambda y: f"=IFERROR({YC[y]}{{np}}/{YC[y]}{{rev}},0)", fcst=lambda c: f"=IFERROR({L('np',c)}/{L('rev',c)},0)", fmt=PCT)
m.row('eps', 'EPS（按期末流通股本，扣除库存股，元）', '元', hist=lambda y: f"=IFERROR({YC[y]}{{np}}/'{S_SH}'!{YC[y]}{REG[(S_SH,'out')]},0)", fcst=lambda c: f"=IFERROR({L('np',c)}/'{S_SH}'!{c}{REG[(S_SH,'out')]},0)", fmt=PS, bold=True)
m.row('eps_w', 'EPS（加权平均股本，元；历史=报告基本每股收益）', '元', hist=lambda y: f"={RAW(HL,'基本每股收益',y)}", fcst=lambda c: f"=IFERROR({L('np',c)}/'{S_SH}'!{c}{REG[(S_SH,'wavg')]},0)", fmt=PS)
m.row('ebit', 'EBIT（剔除利息收支与理财收益；汇兑损益与手续费保留）', hist=lambda y: f"={YC[y]}{{op}}+{YC[y]}{{fin}}-{G(S_FIN,'fx_fee',YC[y])}-{YC[y]}{{inv}}", fcst=lambda c: f"={L('op',c)}+{L('fin',c)}-{G(S_FIN,'fx_fee',c)}-{L('inv',c)}", bold=True)
m.row('da', '折旧与摊销', hist=lambda y: f"={G(S_CAP,'da_tot',YC[y])}", fcst=lambda c: f"={G(S_CAP,'da_tot',c)}")
m.row('ebitda', 'EBITDA', hist=lambda y: f"={YC[y]}{{ebit}}+{YC[y]}{{da}}", fcst=lambda c: f"={L('ebit',c)}+{L('da',c)}", bold=True)
m.row('ebitda_m', '  EBITDA利润率', '%', hist=lambda y: f"=IFERROR({YC[y]}{{ebitda}}/{YC[y]}{{rev}},0)", fcst=lambda c: f"=IFERROR({L('ebitda',c)}/{L('rev',c)},0)", fmt=PCT)
m.row('op_chk', '  校验：营业利润-报告值（历史）', hist=lambda y: f"=ROUND({YC[y]}{{op}}-{RAW(HL,'营业利润',y)},2)")
m.sec('资产负债表')
BSA = [('cash', '货币资金', ['货币资金'], None), ('tfa', '交易性金融资产', ['交易性金融资产'], (S_FIN, 'tfa')),
       ('nrec', '应收票据', ['应收票据'], (S_WC, 'nrec')), ('ar', '应收账款', ['应收账款'], (S_WC, 'ar')), ('rfin', '应收款项融资', ['应收款项融资'], (S_WC, 'rfin')),
       ('pre', '预付款项', ['预付款项'], (S_WC, 'pre')), ('orec', '其他应收款', ['其他应收款(合计)'], (S_WC, 'orec')), ('invt', '存货', ['存货'], (S_WC, 'inv')),
       ('oca', '其他流动资产', ['其他流动资产'], (S_WC, 'oca'))]
for k, nm, labs, link in BSA:
    m.row(k, nm, hist=lambda y, labs=labs: RAWS(HB, labs, y), fcst=(lambda c, link=link: f"={G(link[0],link[1],c)}") if link else (lambda c: f"={L('cash',prev(c))}+{L('cf_net',c)}"), bold=(k == 'cash'))
m.row('ca', '流动资产合计', hist=lambda y: f"=SUM({YC[y]}{{cash}}:{YC[y]}{{oca}})", fcst=lambda c: f"=SUM({L('cash',c)}:{L('oca',c)})", bold=True)
BSN = [('ppe', '固定资产+在建工程', ['固定资产(合计)', '在建工程(合计)'], (S_CAP, 'ppe')), ('intg', '无形资产及长期待摊费用', ['无形资产', '长期待摊费用'], (S_CAP, 'int_close')),
       ('rou', '使用权资产', ['使用权资产'], (S_CAP, 'rou')), ('olta', '商誉、投资性房地产及股权投资', ['长期股权投资', '其他权益工具投资', '投资性房地产', '商誉'], (S_CAP, 'olta_flat')),
       ('dta', '递延所得税资产', ['递延所得税资产'], None), ('onca', '其他非流动资产（含定期存款）', ['其他非流动资产'], (S_FIN, 'onca'))]
for k, nm, labs, link in BSN:
    m.row(k, nm, hist=lambda y, labs=labs: RAWS(HB, labs, y), fcst=(lambda c, link=link: f"={G(link[0],link[1],c)}") if link else (lambda c, k=k: f"={L(k,prev(c))}"))
m.row('nca', '非流动资产合计', hist=lambda y: f"=SUM({YC[y]}{{ppe}}:{YC[y]}{{onca}})", fcst=lambda c: f"=SUM({L('ppe',c)}:{L('onca',c)})", bold=True)
m.row('ta', '资产总计', hist=lambda y: f"={YC[y]}{{ca}}+{YC[y]}{{nca}}", fcst=lambda c: f"={L('ca',c)}+{L('nca',c)}", bold=True)
BSL = [('st_debt', '短期借款', ['短期借款'], (S_FIN, 'st_debt')), ('lt_debt', '长期借款（含一年内到期）', ['长期借款', '一年内到期的非流动负债'], (S_FIN, 'lt_debt')),
       ('npay', '应付票据', ['应付票据'], (S_WC, 'npay')), ('ap', '应付账款', ['应付账款'], (S_WC, 'ap')), ('contract', '合同负债及预收款项', ['预收款项', '合同负债'], (S_WC, 'contract')),
       ('payroll', '应付职工薪酬', ['应付职工薪酬'], (S_WC, 'payroll')), ('taxpay', '应交税费', ['应交税费'], (S_WC, 'taxpay')), ('othpay', '其他应付款', ['其他应付款(合计)'], (S_WC, 'othpay')),
       ('ocl', '其他流动负债', ['其他流动负债'], (S_WC, 'ocl')), ('lease', '租赁负债', ['租赁负债'], (S_FIN, 'lease')),
       ('oncl', '其他非流动负债（递延收益、递延所得税负债、预计负债等）', ['长期应付款(合计)', '预计负债', '递延所得税负债', '递延收益-非流动负债', '其他非流动负债'], None)]
for k, nm, labs, link in BSL:
    m.row(k, nm, hist=lambda y, labs=labs: RAWS(HB, labs, y), fcst=(lambda c, link=link: f"={G(link[0],link[1],c)}") if link else (lambda c, k=k: f"={L(k,prev(c))}"))
m.row('tl', '负债合计', hist=lambda y: f"=SUM({YC[y]}{{st_debt}}:{YC[y]}{{oncl}})", fcst=lambda c: f"=SUM({L('st_debt',c)}:{L('oncl',c)})", bold=True)
m.row('eq_p', '归母所有者权益', hist=lambda y: f"={RAW(HB,'归属于母公司所有者权益合计',y)}", fcst=lambda c: f"={G(S_FIN,'eq_p',c)}")
m.row('eq_m', '少数股东权益', hist=lambda y: f"=N({RAW(HB,'少数股东权益',y)})", fcst=lambda c: f"={L('eq_m',prev(c))}+{L('mi',c)}")
m.row('te', '所有者权益合计', hist=lambda y: f"={YC[y]}{{eq_p}}+{YC[y]}{{eq_m}}", fcst=lambda c: f"={L('eq_p',c)}+{L('eq_m',c)}", bold=True)
m.row('tle', '负债和所有者权益总计', hist=lambda y: f"={YC[y]}{{tl}}+{YC[y]}{{te}}", fcst=lambda c: f"={L('tl',c)}+{L('te',c)}", bold=True)
m.row('bs_chk', '  校验：资产-负债-权益', hist=lambda y: f"=ROUND({YC[y]}{{ta}}-{YC[y]}{{tle}},3)", fcst=lambda c: f"=ROUND({L('ta',c)}-{L('tle',c)},3)")
m.row('ta_chk', '  校验：资产总计-报告值（历史）', hist=lambda y: f"=ROUND({YC[y]}{{ta}}-{RAW(HB,'资产总计',y)},3)")
m.row('tl_chk', '  校验：负债合计-报告值（历史）', hist=lambda y: f"=ROUND({YC[y]}{{tl}}-{RAW(HB,'负债合计',y)},3)")
m.sec('现金流量表（预测为间接法；历史为报告值）')
m.row('cf_ni', '净利润', hist=lambda y: f"={RAW(HCF,'净利润',y)}", fcst=lambda c: f"={L('ni',c)}")
m.row('cf_da', '折旧与摊销', hist=lambda y: RAWS(HCF, ['固定资产折旧、油气资产折耗、生产性生物资产折旧', '无形资产摊销', '使用权资产折旧', '长期待摊费用摊销'], y), fcst=lambda c: f"={L('da',c)}")
m.row('cf_wc', '营运资本变动（增加为负）', hist=lambda y: RAWS(HCF, ['存货的减少', '经营性应收项目的减少', '经营性应付项目的增加'], y), fcst=lambda c: f"=-{G(S_WC,'dnwc',c)}")
m.row('cf_oth', '其他调整（预测：扣除理财收益、加回固定资产减值）', hist=lambda y: f"={RAW(HCF,'经营活动产生的现金流量净额',y)}-{YC[y]}{{cf_ni}}-{YC[y]}{{cf_da}}-{YC[y]}{{cf_wc}}",
      fcst=lambda c: f"=-{L('inv',c)}-{G(S_TAX,'imp_fa',c)}")
m.row('cfo', '经营活动现金流量净额', hist=lambda y: f"={RAW(HCF,'经营活动产生的现金流量净额',y)}", fcst=lambda c: f"={L('cf_ni',c)}+{L('cf_da',c)}+{L('cf_wc',c)}+{L('cf_oth',c)}", bold=True)
m.row('capex', '资本开支（流出为负）', hist=lambda y: f"=-{RAW(HCF,'购建固定资产、无形资产和其他长期资产支付的现金',y)}", fcst=lambda c: f"=-{G(S_CAP,'capex_tot',c)}")
m.row('cf_disp', '处置长期资产收回现金', hist=lambda y: f"=N({RAW(HCF,'处置固定资产、无形资产和其他长期资产收回的现金净额',y)})", fcst=lambda c: f"={G(S_CAP,'disp_nbv',c)}")
m.row('cf_invest', '理财收益、理财与定期存款净额及其他投资', hist=lambda y: f"={RAW(HCF,'投资活动产生的现金流量净额',y)}-{YC[y]}{{capex}}-{YC[y]}{{cf_disp}}",
      fcst=lambda c: f"={L('inv',c)}-({L('tfa',c)}-{L('tfa',prev(c))})-({L('onca',c)}-{L('onca',prev(c))})-({L('dta',c)}-{L('dta',prev(c))})-({L('olta',c)}-{L('olta',prev(c))})")
m.row('cfi', '投资活动现金流量净额', hist=lambda y: f"={RAW(HCF,'投资活动产生的现金流量净额',y)}", fcst=lambda c: f"={L('capex',c)}+{L('cf_disp',c)}+{L('cf_invest',c)}", bold=True)
m.row('cf_debt', '借款净增加', hist=lambda y: f"=N({RAW(HCF,'取得借款收到的现金',y)})-N({RAW(HCF,'偿还债务支付的现金',y)})", fcst=lambda c: f"={L('st_debt',c)}-{L('st_debt',prev(c))}+{L('lt_debt',c)}-{L('lt_debt',prev(c))}")
m.row('cf_div', '现金股利（历史含利息支付）', hist=lambda y: f"=-{RAW(HCF,'分配股利、利润或偿付利息支付的现金',y)}", fcst=lambda c: f"=-{G(S_FIN,'div',c)}")
m.row('cf_bb', '股份回购', hist=lambda y: ("=0" if y != '2025A' else f"=-'{S_SH}'!G{REG[(S_SH,'bb_amt')]}"), fcst=lambda c: f"=-{G(S_FIN,'bb',c)}")
m.row('cf_h', 'H股募资净额及其他股权融资', hist=lambda y: f"=N({RAW(HCF,'吸收投资收到的现金',y)})", fcst=lambda c: f"='{S_SH}'!{c}{REG[(S_SH,'h_net')]}")
m.row('cf_lease', '租赁付款（本金）', fcst=lambda c: f"=-{G(S_CAP,'rou_dep',c)}")
m.row('cf_fo', '其他筹资活动', hist=lambda y: f"={RAW(HCF,'筹资活动产生的现金流量净额',y)}-{YC[y]}{{cf_debt}}-{YC[y]}{{cf_div}}-{YC[y]}{{cf_bb}}-{YC[y]}{{cf_h}}",
      fcst=lambda c: f"=({L('oncl',c)}-{L('oncl',prev(c))})+({L('lease',c)}-{L('lease',prev(c))})")
m.row('cff', '筹资活动现金流量净额', hist=lambda y: f"={RAW(HCF,'筹资活动产生的现金流量净额',y)}", fcst=lambda c: f"={L('cf_debt',c)}+{L('cf_div',c)}+{L('cf_bb',c)}+{L('cf_h',c)}+{L('cf_lease',c)}+{L('cf_fo',c)}", bold=True)
m.row('cf_fx', '汇率变动影响', hist=lambda y: f"={RAW(HCF,'汇率变动对现金的影响',y)}", fcst=lambda c: "=0")
m.row('cf_net', '现金净增加额', hist=lambda y: f"={YC[y]}{{cfo}}+{YC[y]}{{cfi}}+{YC[y]}{{cff}}+{YC[y]}{{cf_fx}}", fcst=lambda c: f"={L('cfo',c)}+{L('cfi',c)}+{L('cff',c)}+{L('cf_fx',c)}", bold=True)
m.row('fcf', '自由现金流（经营现金流-资本开支）', hist=lambda y: f"={YC[y]}{{cfo}}+{YC[y]}{{capex}}", fcst=lambda c: f"={L('cfo',c)}+{L('capex',c)}", bold=True)
m.row('cf_chk', '  校验：现金净增加额-报告值（历史）', hist=lambda y: f"=ROUND({YC[y]}{{cf_net}}-{RAW(HCF,'现金及现金等价物净增加额',y)},2)")
m.row('cash_chk', '  校验：货币资金变动-现金净增加额（预测）', fcst=lambda c: f"=ROUND(({L('cash',c)}-{L('cash',prev(c))})-{L('cf_net',c)},4)")

# ---------------------------------------------------------------- 回报分析
rt = Sheet(S_RET, '回报、杜邦分解与现金转换', 'ROE=净利率×总资产周转率×权益乘数（平均口径）；ROIC=NOPAT/平均投入资本（权益+有息负债+租赁负债-现金类资产）。')
rt.row('npm', '归母净利率', '%', hist=lambda y: f"={G(S_M,'npm',YC[y])}", fcst=lambda c: f"={G(S_M,'npm',c)}", fmt=PCT)
rt.row('at', '总资产周转率（收入/平均总资产）', 'x', hist=lambda y: (f"=IFERROR({G(S_M,'rev',YC[y])}/AVERAGE({G(S_M,'ta',prev(YC[y]))},{G(S_M,'ta',YC[y])}),0)" if y != '2021A' else None), fcst=lambda c: f"=IFERROR({G(S_M,'rev',c)}/AVERAGE({G(S_M,'ta',prev(c))},{G(S_M,'ta',c)}),0)", fmt='0.00"x"')
rt.row('em', '权益乘数（平均总资产/平均归母权益）', 'x', hist=lambda y: (f"=IFERROR(AVERAGE({G(S_M,'ta',prev(YC[y]))},{G(S_M,'ta',YC[y])})/AVERAGE({G(S_M,'eq_p',prev(YC[y]))},{G(S_M,'eq_p',YC[y])}),0)" if y != '2021A' else None), fcst=lambda c: f"=IFERROR(AVERAGE({G(S_M,'ta',prev(c))},{G(S_M,'ta',c)})/AVERAGE({G(S_M,'eq_p',prev(c))},{G(S_M,'eq_p',c)}),0)", fmt='0.00"x"')
rt.row('roe', 'ROE（杜邦，平均口径）', '%', hist=lambda y: (f"={YC[y]}{{npm}}*{YC[y]}{{at}}*{YC[y]}{{em}}" if y != '2021A' else None), fcst=lambda c: f"={L('npm',c)}*{L('at',c)}*{L('em',c)}", fmt=PCT, bold=True)
rt.row('nopat', 'NOPAT（EBIT×(1-实际税率)）', '百万元', hist=lambda y: f"={G(S_M,'ebit',YC[y])}*(1-{G(S_TAX,'etr',YC[y])})", fcst=lambda c: f"={G(S_M,'ebit',c)}*(1-{G(S_TAX,'etr',c)})")
rt.row('ic', '投入资本（权益+有息负债+租赁-现金-理财-其他非流动资产）', '百万元', hist=lambda y: f"={G(S_M,'te',YC[y])}+{G(S_M,'st_debt',YC[y])}+{G(S_M,'lt_debt',YC[y])}+{G(S_M,'lease',YC[y])}-{G(S_M,'cash',YC[y])}-{G(S_M,'tfa',YC[y])}-{G(S_M,'onca',YC[y])}",
       fcst=lambda c: f"={G(S_M,'te',c)}+{G(S_M,'st_debt',c)}+{G(S_M,'lt_debt',c)}+{G(S_M,'lease',c)}-{G(S_M,'cash',c)}-{G(S_M,'tfa',c)}-{G(S_M,'onca',c)}")
rt.row('roic', 'ROIC（NOPAT/平均投入资本）', '%', hist=lambda y: (f"=IFERROR({YC[y]}{{nopat}}/AVERAGE({prev(YC[y])}{{ic}},{YC[y]}{{ic}}),0)" if y != '2021A' else None), fcst=lambda c: f"=IFERROR({L('nopat',c)}/AVERAGE({L('ic',prev(c))},{L('ic',c)}),0)", fmt=PCT, bold=True)
rt.row('iroic', '增量ROIC（ΔNOPAT/Δ投入资本）', '%', hist=lambda y: (f"=IFERROR(({YC[y]}{{nopat}}-{prev(YC[y])}{{nopat}})/({YC[y]}{{ic}}-{prev(YC[y])}{{ic}}),0)" if y != '2021A' else None), fcst=lambda c: f"=IFERROR(({L('nopat',c)}-{L('nopat',prev(c))})/({L('ic',c)}-{L('ic',prev(c))}),0)", fmt=PCT)
rt.row('cconv', '现金转换率（经营现金流/净利润）', '%', hist=lambda y: f"=IFERROR({G(S_M,'cfo',YC[y])}/{G(S_M,'ni',YC[y])},0)", fcst=lambda c: f"=IFERROR({G(S_M,'cfo',c)}/{G(S_M,'ni',c)},0)", fmt=PCT)
rt.row('fcfconv', '自由现金流/净利润', '%', hist=lambda y: f"=IFERROR({G(S_M,'fcf',YC[y])}/{G(S_M,'ni',YC[y])},0)", fcst=lambda c: f"=IFERROR({G(S_M,'fcf',c)}/{G(S_M,'ni',c)},0)", fmt=PCT)
rt.row('fcf_yield', '自由现金流收益率（FCF/当前市值）', '%', fcst=lambda c: f"=IFERROR({G(S_M,'fcf',c)}/({A1('price')}*'{S_SH}'!{c}{REG[(S_SH,'out')]}),0)", fmt=PCT2)
rt.row('pe', 'PE（当前股价/EPS）', 'x', hist=lambda y: f"=IFERROR({A1('price')}/{G(S_M,'eps',YC[y])},0)", fcst=lambda c: f"=IFERROR({A1('price')}/{G(S_M,'eps',c)},0)", fmt=MULT)
rt.row('pb', 'PB（当前股价/每股净资产）', 'x', hist=lambda y: f"=IFERROR({A1('price')}/({G(S_M,'eq_p',YC[y])}/'{S_SH}'!{YC[y]}{REG[(S_SH,'out')]}),0)", fcst=lambda c: f"=IFERROR({A1('price')}/({G(S_M,'eq_p',c)}/'{S_SH}'!{c}{REG[(S_SH,'out')]}),0)", fmt=MULT)
rt.row('ev_ebitda', 'EV/EBITDA（EV=当前股价×2026E年末流通股-2026E年末净现金）', 'x', fcst=lambda c: f"=IFERROR(({A1('price')}*'{S_SH}'!$H${REG[(S_SH,'out')]}-{G(S_FIN,'netcash','H')})/{G(S_M,'ebitda',c)},0)", fmt=MULT)

SCHED = [rv, co, op, cp, wcs, fn, tx, sh, m, rt]
for s in SCHED:
    s.layout()

build_raw(HL, '三环集团[300408.SZ]-利润表.xlsx', '三环集团 历史利润表（合并）')
build_raw(HB, '三环集团[300408.SZ]-资产负债表.xlsx', '三环集团 历史资产负债表（合并）')
build_raw(HCF, '三环集团[300408.SZ]-现金流量表.xlsx', '三环集团 历史现金流量表（合并）')

# ============================================================ 假设 (procedural; registers keys)
from assumptions_v2 import build_assumptions
build_assumptions(wb, S_REV, S_COST, S_OPEX, S_CAP, S_WC, S_FIN, S_TAX, S_M)

# fix placeholder tokens {key} in hist formulas at write time
_orig_write = Sheet.write
def _write(self, wb):
    ws = _orig_write(self, wb)
    import re
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('=') and ('{' in c.value or '<' in c.value):
                def rep(mo):
                    k = mo.group(1)
                    if k == 'r':
                        return str(c.row)
                    return str(REG[(self.name, k)])
                v = re.sub(r'\{([a-z0-9_]+)\}', rep, c.value)
                v = re.sub(r'<([a-z0-9_]+)>', rep, v)
                c.value = v
    return ws
Sheet.write = _write
for s in SCHED:
    s.write(wb)

# ============================================================ remaining sheets

from extras_v2 import build_extras
build_extras(wb, globals())

import os
if os.environ.get('KEEP_SHEETS'):          # render-only copies for visual checks
    keep = os.environ['KEEP_SHEETS'].split(',')
    for w in list(wb.worksheets):
        if w.title not in keep:
            wb.remove(w)
wb.save(OUT)
json.dump({f'{s}|{k}': r for (s, k), r in REG.items()}, open(OUT.rsplit('.', 1)[0] + '_reg.json', 'w'), ensure_ascii=False)
import hist_detail
if hist_detail.MISSING:
    print('WARNING missing note lookups:', hist_detail.MISSING)
print('saved', OUT)
