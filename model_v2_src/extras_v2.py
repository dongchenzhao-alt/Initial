# -*- coding: utf-8 -*-
"""Valuation, sensitivity placeholders, reference and check sheets for the v2 model."""
import json, os
import openpyxl
from openpyxl.styles import Font, Alignment
from engine import *

HERE = os.path.dirname(os.path.abspath(__file__))


def build_extras(wb, g):
    S_REV, S_COST, S_OPEX, S_CAP, S_WC, S_FIN, S_TAX, S_M, S_SH, S_RET = (g[k] for k in (
        'S_REV', 'S_COST', 'S_OPEX', 'S_CAP', 'S_WC', 'S_FIN', 'S_TAX', 'S_M', 'S_SH', 'S_RET'))
    A, A1, SRC = g['A'], g['A1'], g['SRC']

    def MM(k, c):
        return X(S_M, k, c)

    # ================================================================ DCF
    wd = wb.create_sheet('DCF估值'); CUR[0] = 'DCF估值'
    title(wd, 'DCF估值（FCFF，三阶段）与反向DCF',
          '估值基准日2026-12-31（使用2026E年末净现金，含H股募资与定期存款）。第一阶段2027E-2030E取自三表；第二阶段2031E-2035E NOPAT按线性递减增速增长，FCFF=NOPAT×(1-增速÷RONIC)；永续期同理。EBIT税率取各年实际税率。')
    st = {'r': 4}
    DV = {}

    def dl(key, label, f, fmt, fill=None, note_txt=None):
        r = st['r']
        put(wd, f'A{r}', label); put(wd, f'C{r}', f, fmt, fill=fill)
        if note_txt:
            put(wd, f'D{r}', note_txt, italic=True, color='595959')
        DV[key] = f'$C${r}'; REG[('DCF', key)] = r
        st['r'] += 1

    section(wd, st['r'], 'WACC', 14); st['r'] += 1
    dl('ke', '股权成本 = 无风险利率 + Beta × 股权风险溢价', f"={A1('rf')}+{A1('beta')}*{A1('erp')}", PCT)
    dl('kd', '税后债务成本', f"={A1('kd')}*(1-{A1('stat_rate')})", PCT)
    dl('wacc', 'WACC', f"=(1-{A1('wd')})*{DV['ke']}+{A1('wd')}*{DV['kd']}", PCT, fill=KEY_FILL)
    dl('g', '永续增长率', f"={A1('g')}", PCT)
    dl('ronic', '新增投资资本回报率（RONIC）', f"={A1('ronic')}", PCT)
    dl('rev30g', '2030E营业收入增速（第一阶段末）', f"={MM('rev_g','L')}", PCT)
    dl('g2s_eff', '第二阶段起点增速（=MAX(永续增速, 2030E收入增速)+溢价）', f"=MAX({DV['g']},{DV['rev30g']})+{A1('g2s')}", PCT)
    dl('g2e_eff', '第二阶段终点增速（不高于起点、不低于永续增速）', f"=MAX({DV['g']},MIN({A1('g2e')},{DV['g2s_eff']}))", PCT)
    st['r'] += 1
    section(wd, st['r'], '自由现金流（FCFF）', 14); st['r'] += 1
    DY = ['2027E', '2028E', '2029E', '2030E', '2031E', '2032E', '2033E', '2034E', '2035E']
    DC = {y: CL(3 + i) for i, y in enumerate(DY)}
    MCOL = {'2027E': 'I', '2028E': 'J', '2029E': 'K', '2030E': 'L'}
    header_row(wd, st['r'], ['项目（百万元）', ''] + DY); st['r'] += 1
    rowd = {}

    def drow(key, label, fn, fmt=NUM, bold=False):
        r = st['r']
        put(wd, f'A{r}', label, bold=bold)
        for i, y in enumerate(DY):
            f = fn(y, i, DC[y])
            if f is not None:
                put(wd, f'{DC[y]}{r}', f, fmt, bold=bold)
        rowd[key] = r; REG[('DCF', 'row_' + key)] = r
        st['r'] += 1

    drow('t', '折现期（年）', lambda y, i, c: i + 1, '0')
    drow('ebit', 'EBIT（剔除利息收入与理财收益）', lambda y, i, c: f"={MM('ebit', MCOL[y])}" if y in MCOL else None)
    drow('etr', '实际税率', lambda y, i, c: f"={X(S_TAX,'etr',MCOL[y])}" if y in MCOL else None, PCT)
    drow('tax', '减：EBIT所得税', lambda y, i, c: f"=-{c}{rowd['ebit']}*{c}{rowd['etr']}" if y in MCOL else None)
    drow('da', '加：折旧与摊销', lambda y, i, c: f"={MM('da', MCOL[y])}" if y in MCOL else None)
    drow('capex', '减：资本开支', lambda y, i, c: f"={MM('capex', MCOL[y])}" if y in MCOL else None)
    drow('nwc', '减：营运资本增加', lambda y, i, c: f"=-{X(S_WC,'dnwc',MCOL[y])}" if y in MCOL else None)
    drow('disp', '加：处置长期资产收回现金', lambda y, i, c: f"={MM('cf_disp', MCOL[y])}" if y in MCOL else None)
    drow('g2', '增速（第二阶段，线性递减）', lambda y, i, c: None if y in MCOL else f"={DV['g2s_eff']}+({DV['g2e_eff']}-{DV['g2s_eff']})*{i-4}/4", PCT)
    drow('nopat', 'NOPAT（EBIT×(1-税率)；第二阶段按增速滚动）', lambda y, i, c: (f"={c}{rowd['ebit']}+{c}{rowd['tax']}" if y in MCOL
                                                              else f"={prev(c)}{st['r']}*(1+{c}{rowd['g2']})"))
    drow('reinv', '再投资率（第一阶段=1-FCFF/NOPAT；第二阶段=增速÷RONIC）', lambda y, i, c: None, PCT)
    drow('fcff', 'FCFF', lambda y, i, c: (f"={c}{rowd['ebit']}+{c}{rowd['tax']}+{c}{rowd['da']}+{c}{rowd['capex']}+{c}{rowd['nwc']}+{c}{rowd['disp']}" if y in MCOL
                                        else f"={c}{rowd['nopat']}*(1-{c}{rowd['g2']}/{DV['ronic']})"), NUM, True)
    for i, y in enumerate(DY):
        cc = DC[y]
        put(wd, f"{cc}{rowd['reinv']}", (f"=IFERROR(1-{cc}{rowd['fcff']}/{cc}{rowd['nopat']},0)" if y in MCOL else f"={cc}{rowd['g2']}/{DV['ronic']}"), PCT)
    drow('df', '折现系数', lambda y, i, c: f"=1/(1+{DV['wacc']})^{c}{rowd['t']}", '0.0000')
    drow('pv', 'FCFF现值', lambda y, i, c: f"={c}{rowd['fcff']}*{c}{rowd['df']}")
    st['r'] += 1
    section(wd, st['r'], '企业价值与股权价值', 14); st['r'] += 1
    last = DC[DY[-1]]
    dl('sum_pv', 'FCFF现值合计（2027E-2035E）', f"=SUM(C{rowd['pv']}:{last}{rowd['pv']})", NUM)
    dl('tv', '终值（2035E；=NOPAT×(1+g)×(1-g/RONIC)/(WACC-g)）', f"={last}{rowd['nopat']}*(1+{DV['g']})*(1-{DV['g']}/{DV['ronic']})/({DV['wacc']}-{DV['g']})", NUM)
    dl('pv_tv', '终值现值', f"={DV['tv']}*{last}{rowd['df']}", NUM)
    dl('ev', '企业价值（EV）', f"={DV['sum_pv']}+{DV['pv_tv']}", NUM)
    dl('tv_share', '  终值占EV比例', f"=IFERROR({DV['pv_tv']}/{DV['ev']},0)", PCT)
    dl('netcash', '加：2026E年末净现金（货币资金+理财+定期存款-有息负债-租赁负债）', f"={X(S_FIN,'netcash','H')}", NUM)
    dl('mi', '减：少数股东权益', f"={MM('eq_m','H')}", NUM)
    dl('eqv', '股权价值', f"={DV['ev']}+{DV['netcash']}-{DV['mi']}", NUM)
    dl('shares', '总股本（百万股，2026E年末，含库存股，全面摊薄）', f"={X(S_SH,'total','H')}", NUM,
       note_txt='库存股拟用于股权激励/员工持股或维护公司价值，未注销，按全面摊薄计入')
    dl('vps', '每股价值（元，2026-12-31）', f"=IFERROR({DV['eqv']}/{DV['shares']},0)", PS, fill=KEY_FILL)
    dl('vps_now', '每股价值折回2026-09-30（按股权成本折现）', f"={DV['vps']}/(1+{DV['ke']})^{A1('val_lag')}", PS)
    dl('px', '当前A股股价（元）', f"={A1('price')}", PS)
    dl('up', '相对当前股价空间', f"=IFERROR({DV['vps_now']}/{DV['px']}-1,0)", PCT)
    dl('ev_ebitda', '隐含EV/EBITDA（2027E）', f"=IFERROR({DV['ev']}/{MM('ebitda','I')},0)", MULT)
    dl('pe_imp', '隐含PE（2027E，股权价值÷归母净利润）', f"=IFERROR({DV['eqv']}/{MM('np','I')},0)", MULT)
    st['r'] += 1
    section(wd, st['r'], '反向DCF：当前股价隐含了什么（实时公式；其余参数取当前假设）', 14); st['r'] += 1
    dl('mkt_eq', '当前股价对应的2026-12-31股权价值（按股权成本滚动至估值基准日）', f"={A1('price')}*(1+{DV['ke']})^{A1('val_lag')}*{DV['shares']}", NUM)
    dl('mkt_ev', '  减净现金、加少数股东权益后的隐含EV', f"={DV['mkt_eq']}-{DV['netcash']}+{DV['mi']}", NUM)
    dl('tvn', '  隐含终值（2035E，未折现）', f"=({DV['mkt_ev']}-{DV['sum_pv']})/{last}{rowd['df']}", NUM)
    N35 = f"{last}{rowd['nopat']}"
    qa = f"(-{N35}/{DV['ronic']})"; qb = f"({N35}*(1-1/{DV['ronic']})+{DV['tvn']})"; qc = f"({N35}-{DV['tvn']}*{DV['wacc']})"
    dl('g_imp', '隐含永续增长率（WACC、RONIC与前两阶段不变；解二次方程）', f"=IFERROR((-{qb}+SQRT({qb}^2-4*{qa}*{qc}))/(2*{qa}),0)", PCT2, fill=KEY_FILL,
       note_txt='若接近或超过WACC，说明当前股价需要显著更高的前期现金流')
    dl('tv_mult', '  隐含终值/2035E NOPAT', f"=IFERROR({DV['tvn']}/{N35},0)", MULT)
    dl('mkt_ev_ebitda', '当前股价隐含EV/EBITDA（2027E，估值基准日口径）', f"=IFERROR({DV['mkt_ev']}/{MM('ebitda','I')},0)", MULT)
    dl('mkt_pe27', '当前股价对应PE（2027E，流通股口径，与回报分析一致）', f"=IFERROR({A1('price')}/{MM('eps','I')},0)", MULT)
    dl('cash_ps', '  其中：每股净现金（元，2026E年末）', f"=IFERROR({DV['netcash']}/{DV['shares']},0)", PS)
    dl('rev_marker', '隐含WACC与第二阶段增速：见“敏感性分析”表（由重算引擎求解）', '', None)
    st['r'] += 1
    section(wd, st['r'], '敏感性：每股价值（元，2026-12-31）— WACC（行）× 永续增长率（列）', 14); st['r'] += 1
    ws_list = [-0.01, -0.005, 0, 0.005, 0.01]; gs_list = [-0.01, -0.005, 0, 0.005, 0.01]
    r = st['r']
    put(wd, f'A{r}', 'WACC \\ 永续增长率', bold=True)
    for j, dg in enumerate(gs_list):
        put(wd, f'{CL(3+j)}{r}', f"={DV['g']}+({dg})", PCT, bold=True)
    hdr_r = r; r += 1
    fc0, fc1 = f"$C${rowd['fcff']}", f"${last}${rowd['fcff']}"
    t0, t1 = f"$C${rowd['t']}", f"${last}${rowd['t']}"
    for dw in ws_list:
        put(wd, f'B{r}', f"={DV['wacc']}+({dw})", PCT, bold=True)
        for j in range(5):
            gcell = f'{CL(3+j)}${hdr_r}'; wcell = f'$B{r}'
            n35 = f"${last}${rowd['nopat']}"
            put(wd, f'{CL(3+j)}{r}', f"=IFERROR((SUMPRODUCT({fc0}:{fc1}/(1+{wcell})^{t0}:{t1})+{n35}*(1+{gcell})*(1-{gcell}/{DV['ronic']})/({wcell}-{gcell})/(1+{wcell})^{t1}"
                                      f"+{DV['netcash']}-{DV['mi']})/{DV['shares']},0)", PS)
        r += 1
    r += 1
    section(wd, r, '敏感性：每股价值（元）— WACC（行）× 第二阶段起点增速（列，以当前起点为中心；终点取当前终点）', 14); r += 1
    put(wd, f'A{r}', 'WACC \\ 第二阶段起点增速', bold=True)
    for j, dg in enumerate([-0.04, -0.02, 0, 0.02, 0.04]):
        put(wd, f'{CL(3+j)}{r}', f"={DV['g2s_eff']}+({dg})", PCT, bold=True)
    hdr2 = r; r += 1
    for dw in ws_list:
        put(wd, f'B{r}', f"={DV['wacc']}+({dw})", PCT, bold=True)
        for j in range(5):
            gs = f'{CL(3+j)}${hdr2}'; w = f'$B{r}'; ge = DV['g2e_eff']; n30 = f"$F${rowd['nopat']}"; ro = DV['ronic']
            gk = [f"({gs}+({ge}-{gs})*{k}/4)" for k in range(5)]
            pv1 = f"SUMPRODUCT($C${rowd['fcff']}:$F${rowd['fcff']}/(1+{w})^$C${rowd['t']}:$F${rowd['t']})"
            cum, pv2 = [], []
            for k in range(5):
                cum.append(f"(1+{gk[k]})")
                pv2.append(f"{n30}*{'*'.join(cum)}*(1-{gk[k]}/{ro})/(1+{w})^{5+k}")
            n35 = f"{n30}*{'*'.join(cum)}"
            put(wd, f'{CL(3+j)}{r}', f"=IFERROR(({pv1}+{'+'.join(pv2)}+{n35}*(1+{DV['g']})*(1-{DV['g']}/{ro})/({w}-{DV['g']})/(1+{w})^9+{DV['netcash']}-{DV['mi']})/{DV['shares']},0)", PS)
        r += 1
    wd.column_dimensions['A'].width = 58; wd.column_dimensions['B'].width = 10
    for i in range(9):
        wd.column_dimensions[CL(3 + i)].width = 12
    wd.sheet_view.showGridLines = False

    def DVX(k):
        return f"'DCF估值'!{DV[k]}"

    # ================================================================ relative valuation
    wr = wb.create_sheet('相对估值'); CUR[0] = '相对估值'
    title(wr, '相对估值（PE/PB）与可比公司', '可比公司倍数请按最新行情填入黄底单元格（模型不预填，避免使用过期数据）。')
    r = 4
    section(wr, r, 'PE估值区间（元/股）', 10); r += 1
    header_row(wr, r, ['PE倍数', ''] + ['2026E', '2027E', '2028E']); r += 1
    for k in range(6):
        put(wr, f'A{r}', f"={A1('pe_lo')}+({A1('pe_hi')}-{A1('pe_lo')})*{k}/5", MULT, bold=True)
        for j, mc in enumerate(['H', 'I', 'J']):
            put(wr, f'{CL(3+j)}{r}', f"=$A{r}*{MM('eps', mc)}", PS)
        r += 1
    for lab, fn, fmt in (('EPS（模型，摊薄）', lambda mc: f"={MM('eps', mc)}", PS),
                         ('当前股价对应PE', lambda mc: f"={X(S_RET,'pe',mc)}", MULT),
                         ('当前股价对应PB', lambda mc: f"={X(S_RET,'pb',mc)}", MULT),
                         ('当前股价对应EV/EBITDA', lambda mc: f"={X(S_RET,'ev_ebitda',mc)}", MULT)):
        put(wr, f'A{r}', lab, bold=True)
        for j, mc in enumerate(['H', 'I', 'J']):
            put(wr, f'{CL(3+j)}{r}', fn(mc), fmt)
        r += 1
    r += 1
    section(wr, r, '可比公司（请填入最新倍数）', 10); r += 1
    header_row(wr, r, ['公司', '代码', 'PE 2026E', 'PE 2027E', 'PB（最新）', '备注']); r += 1
    peers = [('风华高科', '000636.SZ', 'MLCC、电阻'), ('国瓷材料', '300285.SZ', '电子陶瓷粉体、MLCC介质材料'), ('天孚通信', '300394.SZ', '光器件'),
             ('村田制作所', '6981.T', 'MLCC龙头'), ('太阳诱电', '6976.T', 'MLCC'), ('京瓷', '6971.T', '精密陶瓷、电子元件'), ('三星电机', '009150.KS', 'MLCC')]
    p0 = r
    for nm, code, rem in peers:
        put(wr, f'A{r}', nm); put(wr, f'B{r}', code)
        for col in 'CDE':
            wr[f'{col}{r}'].fill = KEY_FILL; wr[f'{col}{r}'].number_format = MULT
            wr[f'{col}{r}'].font = Font(name=F, size=10, color=BLUE)
        put(wr, f'F{r}', rem, italic=True, color='595959')
        r += 1
    put(wr, f'A{r}', '可比公司中位数', bold=True)
    for col in 'CDE':
        put(wr, f'{col}{r}', f"=IFERROR(MEDIAN({col}{p0}:{col}{r-1}),0)", MULT, bold=True)
    med = r; r += 1
    put(wr, f'A{r}', '按可比中位数的三环每股价值（元）', bold=True)
    put(wr, f'C{r}', f"=C{med}*{MM('eps','H')}", PS); put(wr, f'D{r}', f"=D{med}*{MM('eps','I')}", PS)
    put(wr, f'E{r}', f"=E{med}*{X(S_FIN,'eq_p','H')}/{X(S_SH,'out','H')}", PS)
    wr.column_dimensions['A'].width = 34; wr.column_dimensions['B'].width = 12
    for col in 'CDE':
        wr.column_dimensions[col].width = 12
    wr.column_dimensions['F'].width = 30
    wr.sheet_view.showGridLines = False

    # ================================================================ SOFC sensitivity (live)
    wf = wb.create_sheet('SOFC敏感性'); CUR[0] = 'SOFC敏感性'
    title(wf, 'SOFC隔膜片敏感性（2027E，实时公式）', '收入=Bloom电池片层需求（GW）×(1000/单片功率W)×三环份额×单价；需求与单片功率取“假设”当前情景2027E采用值。')
    put(wf, 'A4', '2027E Bloom电池片层需求（GW）'); put(wf, 'C4', f"={A('sofc_gw','I')}", '0.00')
    put(wf, 'A5', '2027E单片功率（W/片）'); put(wf, 'C5', f"={A('sofc_w','I')}", '0.0')
    put(wf, 'A6', 'SOFC隔膜片毛利率（2027E采用值）'); put(wf, 'C6', f"={A('sofc_gm','I')}", PCT)
    put(wf, 'A7', '毛利转净利系数（取2027E公司整体归母净利润÷毛利）'); put(wf, 'C7', f"=IFERROR({MM('np','I')}/{MM('gp','I')},0)", PCT)
    note(wf, 'C7', '以公司整体毛利转化为归母净利润的比例近似SOFC隔膜片的转化率（含其他收益、利息收入等）')
    put(wf, 'A8', '2027E模型归母净利润'); put(wf, 'C8', f"={MM('np','I')}", NUM)
    put(wf, 'A9', '2027E模型SOFC隔膜片收入'); put(wf, 'C9', f"={X(S_REV,'sofc_rev','I')}", NUM)
    r = 11
    prices = [16, 18, 20, 22, 24, 26]; shares = [0.4, 0.5, 0.6, 0.7, 0.8, 0.9]
    for ttl, fn, fmt in (('2027E SOFC隔膜片收入（百万元）：三环份额（行）× 单价（列）', lambda rr, c, h: f"=$C$4*1000/$C$5*$B{rr}*{c}${h}", NUM),
                         ('SOFC隔膜片净利润占2027E归母净利润比例', lambda rr, c, h: f"=IFERROR($C$4*1000/$C$5*$B{rr}*{c}${h}*$C$6*$C$7/$C$8,0)", PCT)):
        section(wf, r, ttl, 10); r += 1
        put(wf, f'A{r}', '份额 \\ 单价（元/片）', bold=True)
        for j, p in enumerate(prices):
            put(wf, f'{CL(3+j)}{r}', p, '0', bold=True)
        h = r; r += 1
        for s in shares:
            put(wf, f'B{r}', s, PCT, bold=True)
            for j in range(6):
                put(wf, f'{CL(3+j)}{r}', fn(r, CL(3 + j), h), fmt)
            r += 1
        r += 1
    section(wf, r, 'Bloom需求（GW，行）× 单片功率（W，列）：2027E SOFC收入（份额与单价取当前情景）', 10); r += 1
    gws = [1.5, 2.0, 2.5, 3.0, 3.5, 4.0]; wts = [25, 27.5, 30, 32.5, 35, 37.5]
    put(wf, f'A{r}', 'GW \\ W', bold=True)
    for j, w in enumerate(wts):
        put(wf, f'{CL(3+j)}{r}', w, '0.0', bold=True)
    h3 = r; r += 1
    for gv in gws:
        put(wf, f'B{r}', gv, '0.0', bold=True)
        for j in range(6):
            put(wf, f'{CL(3+j)}{r}', f"=$B{r}*1000/{CL(3+j)}${h3}*{A('sofc_sh','I')}*{A('sofc_px','I')}", NUM)
        r += 1
    wf.column_dimensions['A'].width = 44; wf.column_dimensions['B'].width = 10
    for j in range(6):
        wf.column_dimensions[CL(3 + j)].width = 11
    wf.sheet_view.showGridLines = False

    # ================================================================ sensitivity + scenario sheets (static results from sens_v2.py)
    from sens_render import render_sens
    render_sens(wb, g, DV)

    # ================================================================ consensus
    wc = wb.create_sheet('一致预期对比'); CUR[0] = '一致预期对比'
    title(wc, '模型与Wind一致预期对比（仅作参照，不作为模型输入）', '资料来源：Wind盈利预测（截至2026-09-30，12家机构）。一致预期EPS隐含股本约1,995-1,999百万股（H股发行后总股本），模型EPS按扣除库存股的期末流通股本，EPS比较以归母净利润为准。')
    header_row(wc, 4, ['项目（百万元）', ''] + ['2026E', '2027E', '2028E'])
    cons = {'营业收入（一致预期）': [13140, 17470, 20690], '归母净利润（一致预期）': [4198, 6142, 7143], 'EPS（一致预期，元）': [2.10, 3.08, 3.58]}
    r = 5; crow = {}
    for lab, vals in cons.items():
        put(wc, f'A{r}', lab)
        for j, v in enumerate(vals):
            put(wc, f'{CL(3+j)}{r}', v, PS if 'EPS' in lab else NUM)
        crow[lab] = r; r += 1
    for lab, k, fmt in (('营业收入（模型）', 'rev', NUM), ('归母净利润（模型）', 'np', NUM), ('EPS（模型，元）', 'eps', PS)):
        put(wc, f'A{r}', lab, bold=True)
        for j, mc in enumerate(['H', 'I', 'J']):
            put(wc, f'{CL(3+j)}{r}', f"={MM(k, mc)}", fmt, bold=True)
        crow[lab] = r; r += 1
    for a, b, lab in (('营业收入（模型）', '营业收入（一致预期）', '营业收入：模型/一致预期-1'), ('归母净利润（模型）', '归母净利润（一致预期）', '归母净利润：模型/一致预期-1'),
                      ('EPS（模型，元）', 'EPS（一致预期，元）', 'EPS：模型/一致预期-1')):
        put(wc, f'A{r}', lab)
        for j in range(3):
            put(wc, f'{CL(3+j)}{r}', f"=IFERROR({CL(3+j)}{crow[a]}/{CL(3+j)}{crow[b]}-1,0)", PCT)
        r += 1
    r += 1
    section(wc, r, '一致预期分布（2026E归母净利润，百万元）', 6); r += 1
    for lab, v in (('中值', 4206), ('最大值', 5503), ('最小值', 3367), ('标准差', 720.8), ('预测机构数', 12)):
        put(wc, f'A{r}', lab); put(wc, f'C{r}', v, NUM); r += 1
    wc.column_dimensions['A'].width = 34
    for col in 'CDE':
        wc.column_dimensions[col].width = 12
    wc.sheet_view.showGridLines = False

    # ================================================================ quarterly
    wsq = wb.create_sheet('季度数据'); CUR[0] = '季度数据'
    title(wsq, '三环集团 单季度利润表、现金流与利润率', '资料来源：Wind单季度利润表与现金流量表（公司公告值），万元÷100=百万元。增速与利润率为公式。')
    qws = openpyxl.load_workbook(SRC + '三环集团[300408.SZ]-利润表(单季).xlsx', data_only=True).active
    qrows = list(qws.iter_rows(values_only=True)); qhdr = qrows[1]
    QCOLS = ['2024第一季度', '2024第二季度', '2024第三季度', '2024第四季度', '2025第一季度', '2025第二季度', '2025第三季度', '2025第四季度', '2026第一季度', '2026第二季度']
    QLAB = ['2024Q1', '2024Q2', '2024Q3', '2024Q4', '2025Q1', '2025Q2', '2025Q3', '2025Q4', '2026Q1', '2026Q2']
    qidx = [qhdr.index(c) for c in QCOLS]
    header_row(wsq, 4, ['项目（百万元）', ''] + QLAB)
    QITEMS = [('营业收入', '营业收入'), ('营业成本', '营业成本'), ('税金及附加', '税金及附加'), ('销售费用', '销售费用'), ('管理费用', '管理费用'),
              ('研发费用', '研发费用'), ('财务费用', '财务费用'), ('其他收益', '加：其他收益'), ('营业利润', '营业利润'), ('利润总额', '利润总额'),
              ('所得税', '减：所得税'), ('净利润', '净利润'), ('归母净利润', '归属于母公司所有者的净利润')]
    qr = {}; r = 5
    for lab, wl in QITEMS:
        src_row = [row for row in qrows if row[0] and str(row[0]).strip() == wl][0]
        put(wsq, f'A{r}', lab, bold=lab in ('营业收入', '归母净利润', '营业利润'))
        for j, i in enumerate(qidx):
            put(wsq, f'{CL(3 + j)}{r}', round(src_row[i] / 100, 4), NUM)
        qr[lab] = r; REG[('季度数据', lab)] = r; r += 1
    qcf = openpyxl.load_workbook(SRC + '三环集团[300408.SZ]-现金流量表(单季).xlsx', data_only=True).active
    qcrows = list(qcf.iter_rows(values_only=True)); qchdr = qcrows[1]
    qcidx = [qchdr.index(c) for c in QCOLS]
    for lab, wl in (('经营活动现金流量净额', '经营活动产生的现金流量净额'), ('购建长期资产支付的现金', '购建固定资产、无形资产和其他长期资产支付的现金'),
                    ('分配股利、利润或偿付利息支付的现金', '分配股利、利润或偿付利息支付的现金')):
        src_row = [row for row in qcrows if row[0] and str(row[0]).strip() == wl][0]
        put(wsq, f'A{r}', lab, bold=lab.startswith('经营'))
        for j, i in enumerate(qcidx):
            v = src_row[i]
            put(wsq, f'{CL(3 + j)}{r}', round((v or 0) / 100, 4), NUM)
        qr[lab] = r; REG[('季度数据', lab)] = r; r += 1
    r += 1
    section(wsq, r, '利润率与增速（公式）', 12); r += 1
    def qrow(lab, fn, fmt=PCT):
        nonlocal r
        put(wsq, f'A{r}', lab)
        for j in range(10):
            f = fn(CL(3 + j), j)
            if f is not None:
                put(wsq, f'{CL(3 + j)}{r}', f, fmt)
        r += 1
    qrow('毛利率', lambda c, j: f'=IFERROR(1-{c}{qr["营业成本"]}/{c}{qr["营业收入"]},0)')
    qrow('期间费用率（销售+管理+研发）', lambda c, j: f'=IFERROR(({c}{qr["销售费用"]}+{c}{qr["管理费用"]}+{c}{qr["研发费用"]})/{c}{qr["营业收入"]},0)')
    qrow('营业利润率', lambda c, j: f'=IFERROR({c}{qr["营业利润"]}/{c}{qr["营业收入"]},0)')
    qrow('归母净利率', lambda c, j: f'=IFERROR({c}{qr["归母净利润"]}/{c}{qr["营业收入"]},0)')
    qrow('所得税实际税率', lambda c, j: f'=IFERROR({c}{qr["所得税"]}/{c}{qr["利润总额"]},0)')
    qrow('营业收入同比', lambda c, j: None if j < 4 else f'=IFERROR({c}{qr["营业收入"]}/{CL(3 + j - 4)}{qr["营业收入"]}-1,0)')
    qrow('营业收入环比', lambda c, j: None if j < 1 else f'=IFERROR({c}{qr["营业收入"]}/{CL(3 + j - 1)}{qr["营业收入"]}-1,0)')
    qrow('归母净利润同比', lambda c, j: None if j < 4 else f'=IFERROR({c}{qr["归母净利润"]}/{CL(3 + j - 4)}{qr["归母净利润"]}-1,0)')
    qrow('归母净利润环比', lambda c, j: None if j < 1 else f'=IFERROR({c}{qr["归母净利润"]}/{CL(3 + j - 1)}{qr["归母净利润"]}-1,0)')
    r += 1
    section(wsq, r, '2026E隐含下半年（模型全年-2026H1实际）', 12); r += 1
    for lab, k in (('营业收入', 'rev'), ('归母净利润', 'np')):
        put(wsq, f'A{r}', f'2026E {lab}（模型）'); put(wsq, f'C{r}', f"={MM(k,'H')}", NUM); r += 1
        put(wsq, f'A{r}', f'  2026H1 {lab}（实际）'); put(wsq, f'C{r}', f"=K{qr[lab]}+L{qr[lab]}", NUM); r += 1
        put(wsq, f'A{r}', f'  隐含2026H2 {lab}'); put(wsq, f'C{r}', f"=C{r-2}-C{r-1}", NUM); r += 1
        put(wsq, f'A{r}', f'  隐含H2/H1'); put(wsq, f'C{r}', f"=IFERROR(C{r-1}/C{r-2},0)", '0.00"x"'); r += 1
        put(wsq, f'A{r}', f'  参照：2025H2/2025H1'); put(wsq, f'C{r}', f"=IFERROR((I{qr[lab]}+J{qr[lab]})/(G{qr[lab]}+H{qr[lab]}),0)", '0.00"x"'); r += 1
    wsq.column_dimensions['A'].width = 34; wsq.column_dimensions['B'].width = 4
    for j in range(10):
        wsq.column_dimensions[CL(3 + j)].width = 11
    wsq.freeze_panes = 'C5'; wsq.sheet_view.showGridLines = False

    # ================================================================ research notes
    wn = wb.create_sheet('调研要点'); CUR[0] = '调研要点'
    title(wn, '产业链调研与公开信息要点（用于SOFC子模型与需求假设）', '来源：Third Bridge访谈纪要（2026-04-03、2026-06-11、2026-07-22）及公开资料；访谈观点为专家个人判断，仅作假设参考。')
    header_row(wn, 4, ['日期', '来源', '要点', '模型中的用途'])
    notes = [
        ('2026-04-03', 'Bloom前高级工程经理', 'Bloom对外产能指电池片与电堆层；其中约25%-30%用于翻新与现场更换；2026年底电池片与电堆产能约2GW，收入端上限约1.5GW', '电池片层需求与验收量的换算'),
        ('2026-04-03', 'Bloom前高级工程经理', '2025年验收约500MW，四季度产能折年约800MW；2026年实际验收约1.25-1.5GW；2027年底收入端产能2.5-3GW；2028-29年电堆产能目标约5GW', 'Bloom电池片层需求（GW）各情景'),
        ('2026-04-03', 'Bloom前高级工程经理', '约两年后推出新一代产品，单位产能对应更多kW；未来两年首次面临电池片与电堆的制造约束', '单片功率提升假设；份额下降时点'),
        ('2026-04-03', 'Bloom前高级工程经理', '系统价格（含激励后）约3,300-3,600美元/kW；部分合约发货即确认验收，现场服务与特种焊接人手是验收瓶颈', '采购与验收时差（库存）'),
        ('2026-06-11', '三环前业务拓展负责人', '与Bloom的隔膜片合作十几年，从氧化铝陶瓷开始；隔膜片市场体量小，三环利用富余产能改造设备切入；CoorsTek具备制造能力，日本厂商成本不占优', '二供进入节奏；第一大客户销售额口径'),
        ('2026-06-11', '三环前业务拓展负责人', '隔膜片对三环整体营收、利润贡献不大；三环同时自研SOFC电堆，已出货', 'SOFC对利润贡献的交叉验证'),
        ('2026-06-11', '三环前业务拓展负责人', '国内MLCC涨价集中在高容产品，低容不涨；AI服务器高端MLCC使用部分稀土材料', '电子元件量价拆分'),
        ('2026-07-22', 'Bloom前国际业务高级管理人员', '65kW约用2,000片电池片（约32.5W/片），650kW约2万片；瓶颈在电池片烧结速度，Bloom已在找第二来源', '单片功率口径；二供'),
        ('2026-07-22', 'Bloom前国际业务高级管理人员', 'SOFC在北美大型数据中心发电中约占10%，2030年前后或升至约20%；数据中心负载波动使反应器温度变化增多', 'Bloom需求上限；翻新需求'),
        ('2026-07-28', 'Bloom二季度业绩公告', '二季度收入10.65亿美元，同比+166%；2026年收入指引上调至39-42亿美元（中值同比约翻倍）', '2026E Bloom需求约为2025年两倍'),
        ('2026-07-29', '首尔经济新闻', 'Amosense已向Bloom交付首批SOFC陶瓷基板，三季度内月产能由20万片扩至60万片', '二供产能（份额假设）'),
        ('2026-08-28', '三环2026年半年报', '营收64.23亿元（+54.82%），归母净利19.32亿元（+56.18%）；增长来自MLCC与插芯、套筒等光通信产品；H股共发行80,221,200股；在建工程升至848.4百万元', '2026E收入基数、股本与资本开支'),
    ]
    r = 5
    for d, s, t, u in notes:
        put(wn, f'A{r}', d); put(wn, f'B{r}', s); put(wn, f'C{r}', t, wrap=True); put(wn, f'D{r}', u, wrap=True)
        wn.row_dimensions[r].height = 42; r += 1
    wn.column_dimensions['A'].width = 12; wn.column_dimensions['B'].width = 26; wn.column_dimensions['C'].width = 90; wn.column_dimensions['D'].width = 34
    wn.sheet_view.showGridLines = False

    # ================================================================ disclosure detail (extracted note tables)
    wx = wb.create_sheet('披露明细'); CUR[0] = '披露明细'
    title(wx, '财务报表附注明细（年报/半年报原文提取，供核对）', '金额单位为元或万元的表已换算为百万元（比例、单价、股数、人数等非金额行列保持原值）；其他单位按原披露列示，以每张表的单位行为准。模型明细表中的历史数据均取自本表或Wind报表。')
    r = 4
    EXT = os.path.join(HERE, 'extract')
    DOMS = [('D1_fixed_assets', '固定资产、在建工程、无形资产与折旧摊销'), ('D2_financing_cash', '货币资金、理财、借款、分红与回购'),
            ('D3_revenue_cost', '收入构成、成本构成、员工与期间费用'), ('D4_working_capital', '应收、存货、应付与其他营运项目'),
            ('D5_tax_other_pl', '所得税、其他收益、减值与非经常性损益'), ('D6_capacity_segments_misc', '产能、分部与其他')]
    for dom, dname in DOMS:
        p = None
        for pref in ('', 'raw_'):
            if os.path.exists(os.path.join(EXT, f'{pref}{dom}.json')):
                p = os.path.join(EXT, f'{pref}{dom}.json'); break
        if not p:
            continue
        data = json.load(open(p))
        section(wx, r, f'{dom[:2]}　{dname}', 12); r += 1
        for t in data['tables']:
            unit = t.get('unit', '')
            u0 = unit.strip()
            is_yuan = u0.startswith('元') and not u0.startswith('元/股')
            is_wan = u0.startswith('万元')
            SCALE = 1e6 if is_yuan else 100.0
            is_yuan = is_yuan or is_wan
            put(wx, f'A{r}', f"{t.get('title', t['id'])}", bold=True, color='1F3864')
            r += 1
            put(wx, f'A{r}', f"单位：{('百万元（原表单位' + ('万元' if is_wan else '元') + '；比例、单价、股数、人数等非金额行列保持原值；原单位说明：' + unit + '）') if is_yuan else unit}；来源：{str(t.get('source',''))[:200]}", italic=True, color='595959')
            r += 1
            cols = t['columns']
            header_row(wx, r, ['项目', ''] + [str(c)[:24] for c in cols], fill=openpyxl.styles.PatternFill('solid', fgColor='44546A')); r += 1
            NONMONEY = ('pct', '_pp', '比例', '占比', '比重', '人数', '人员数量', '（人）', '(人)', '股数', '户数', '（股）', '(股)', '天数', '每股', '元/股', '/10股', '（%）', '(%)')
            for row in t['rows']:
                put(wx, f'A{r}', str(row['label'])[:80])
                nums = [abs(x) for x in row['values'] if isinstance(x, (int, float)) and x]
                small_row = bool(nums) and max(nums) < 1000          # ratios, rates, prices, counts: never amounts in 元
                for j, v in enumerate(row['values']):
                    if isinstance(v, (int, float)):
                        cj = str(cols[j])
                        pct_col = any(k in cj for k in ('pct', '%', '比例', '进度', '（股）', '(股)', '股数', '户', '元/股', '港元', '（人）', '(人)')) or cj.endswith('率')
                        lab_s = str(row['label'])
                        non_money = small_row or any(k in lab_s for k in NONMONEY)
                        if is_yuan and not pct_col and not non_money:
                            put(wx, f'{CL(3+j)}{r}', round(v / SCALE, 4), NUM)
                        else:
                            put(wx, f'{CL(3+j)}{r}', v, '#,##0.00;(#,##0.00);"-"')
                    elif v is not None:
                        put(wx, f'{CL(3+j)}{r}', str(v)[:30])
                r += 1
            r += 1
    wx.column_dimensions['A'].width = 60; wx.column_dimensions['B'].width = 2
    for j in range(24):
        wx.column_dimensions[CL(3 + j)].width = 14
    wx.freeze_panes = 'C4'; wx.sheet_view.showGridLines = False

    # ================================================================ checks
    wk = wb.create_sheet('校验'); CUR[0] = '校验'
    title(wk, '模型校验（全部应为0）', '任何非零值表示勾稽关系出错；汇总在最下方。')
    header_row(wk, 4, ['校验项', '所在表'] + YEARS)
    r = 5
    CHECKS = [('资产负债表配平（资产-负债-权益）', S_M, 'bs_chk', YEARS), ('历史资产总计与报告值差异', S_M, 'ta_chk', HIST),
              ('历史负债合计与报告值差异', S_M, 'tl_chk', HIST), ('历史营业利润重算与报告值差异', S_M, 'op_chk', HIST),
              ('历史现金净增加额与报告值差异', S_M, 'cf_chk', HIST), ('预测货币资金变动与现金净增加额差异', S_M, 'cash_chk', FCST),
              ('固定资产类别合计与资产负债表差异（历史）', S_CAP, 'fa_chk', ['2022A', '2023A', '2024A', '2025A']),
              ('存货分类合计与资产负债表差异（历史）', S_WC, 'inv_chk', ['2023A', '2024A', '2025A']),
              ('权益明细合计与报告值差异（历史）', S_FIN, 'eq_chk', HIST),
              ('营业成本按性质合计与营业成本差异', S_COST, 'nat_chk', ['2023A', '2024A', '2025A'] + FCST)]
    for lab, sh, k, ys in CHECKS:
        put(wk, f'A{r}', lab); put(wk, f'B{r}', sh, color='595959')
        for y in ys:
            put(wk, f'{YC[y]}{r}', f"=ROUND(N({X(sh, k, YC[y])}),2)", NUM)
        r += 1
    put(wk, f'A{r}', '分部收入合计与利润表收入差异'); put(wk, f'B{r}', S_REV, color='595959')
    for y in ['2024A', '2025A'] + FCST:
        c = YC[y]
        put(wk, f'{c}{r}', f"=ROUND({X(S_REV,'rev_tot',c)}-{MM('rev',c)},2)", NUM)
    r += 1
    put(wk, f'A{r}', '分部营业成本合计与利润表营业成本差异'); put(wk, f'B{r}', S_COST, color='595959')
    for y in FCST:
        c = YC[y]
        put(wk, f'{c}{r}', f"=ROUND({X(S_COST,'cogs_tot',c)}-{MM('cogs',c)},2)", NUM)
    r += 1
    put(wk, f'A{r}', '2025A分部营业成本（年报）合计与利润表差异'); put(wk, f'B{r}', S_COST, color='595959')
    c = 'G'
    put(wk, f'{c}{r}', f"=ROUND({X(S_COST,'cogs_ec',c)}+{X(S_COST,'cogs_cd',c)}+{X(S_COST,'cogs_cm',c)}+{X(S_COST,'cogs_eq',c)}+{X(S_COST,'cogs_ot',c)}+{X(S_COST,'cogs_ob',c)}-{MM('cogs',c)},2)", NUM)
    r += 1
    put(wk, f'A{r}', '新增固定资产结构合计-100%（预测）'); put(wk, f'B{r}', '假设', color='595959')
    for c in FCOL:
        put(wk, f'{c}{r}', f"=ROUND({A('mix_bld',c)}+{A('mix_mach',c)}+{A('mix_trans',c)}+{A('mix_elec',c)}-1,6)", NUM)
    r += 1
    put(wk, f'A{r}', '现金为负的年份（预测，1=是）'); put(wk, f'B{r}', S_M, color='595959')
    for c in FCOL:
        put(wk, f'{c}{r}', f"=IF({MM('cash',c)}<0,1,0)", '0')
    r += 1
    r += 1
    put(wk, f'A{r}', '校验汇总（所有校验绝对值之和，应为0）', bold=True)
    put(wk, f'C{r}', f"=SUMPRODUCT(ABS(C5:L{r-2}))", NUM, bold=True, fill=KEY_FILL)
    REG[('校验', 'sum')] = r
    r += 2
    put(wk, f'A{r}', '提示：敏感性冲击是否未归零（1=是；不计入校验汇总）'); put(wk, f'B{r}', '假设', color='595959')
    shock_keys = [k for (s, k) in REG if s == '假设1' and k.startswith('s_')]
    put(wk, f'C{r}', '=IF(' + '+'.join(f'ABS({A1(k)})' for k in shock_keys) + '>0,1,0)', '0')
    REG[('校验', 'shockflag')] = r
    r += 1
    put(wk, f'A{r}', '提示：静态敏感性结果与实时模型不一致（1=需重跑sens_v2.py；不计入校验汇总）'); put(wk, f'B{r}', '敏感性分析', color='595959')
    put(wk, f'C{r}', f"=IF(OR({A1('scen')}<>1,C{REG[('校验','shockflag')]}=1),0,IF(ABS('敏感性分析'!C{REG.get(('敏感性分析','vps'), 9)}-'DCF估值'!{DV['vps']})>0.01,1,0))", '0')
    put(wk, f'D{r}', '仅在基准情景且冲击为0时判断（静态结果为基准情景）', italic=True, color='595959')
    REG[('校验', 'staleflag')] = r
    wk.column_dimensions['A'].width = 52; wk.column_dimensions['B'].width = 14
    for col in YC.values():
        wk.column_dimensions[col].width = 11
    wk.sheet_view.showGridLines = False

    # ================================================================ cover
    wv = wb.create_sheet('封面与摘要'); CUR[0] = '封面与摘要'
    title(wv, '三环集团（300408.SZ / 06951.HK）详细财务与估值模型', '编制日期：2026-09-30｜单位：百万元人民币（另注明除外）｜数据截至2026年半年报（2026-08-28）')
    put(wv, 'A4', '使用说明', bold=True, color='1F3864')
    instr = ['1. 在“假设”C5切换情景（1=基准，2=乐观，3=悲观）；蓝字为录入值/假设，黑字为本表公式，绿字为跨表链接，灰底为预测列，黄底为关键假设。',
             '2. 模型链路：假设 → 收入拆解（量价、SOFC子模型）→ 成本拆解（分部毛利率、成本按性质）→ 人员与费用（人数×人均薪酬）→ 资本开支与折旧（在建工程、四类固定资产）→ 营运资金（分项周转天数）→ 融资与资金（利息、权益变动）→ 税务与其他损益（所得税桥）→ 三表预测 → 回报分析 → DCF/相对估值。',
             '3. 利息收入按期初余额计息、利息费用按平均借款余额，模型无循环引用；货币资金由现金流量表倒挤，资产负债表自动配平。营业成本以分部毛利率为主驱动，生产人工与生产折旧偏离基准标定值的部分直接计入营业成本；DCF第二阶段与永续期按“再投资率=增速÷RONIC”计算FCFF。',
             '4. “假设”第四节为敏感性冲击单元（默认0）；“敏感性分析”“情景对比”两表为重算引擎写入的静态结果，修改假设后需重跑才会更新；DCF表内的WACC×增长率表与反向DCF为实时公式。',
             '5. 历史数据：三张报表来自Wind（公司公告值）；附注明细（固定资产变动、成本构成、员工、存货、税率调节等）取自年报/半年报原文，见“披露明细”。',
             '6. “校验”表汇总应为0；A股股价、汇率、可比公司倍数请按最新行情更新。']
    for i, t in enumerate(instr):
        put(wv, f'A{5+i}', t)
    r = 5 + len(instr) + 1
    section(wv, r, '核心输出（随情景自动更新）', 9); r += 1
    header_row(wv, r, ['指标', '', '2025A', '2026E', '2027E', '2028E', '2029E', '2030E']); r += 1
    outs = [('营业收入', S_M, 'rev', NUM), ('  同比', S_M, 'rev_g', PCT), ('  电子元件', S_REV, 'ec_rev', NUM), ('  通信器件', S_REV, 'cd_rev', NUM),
            ('  电子及陶瓷材料', S_REV, 'cm_rev', NUM), ('    其中SOFC隔膜片（2025A为第一大客户销售额代理）', S_REV, 'sofc_rev', NUM), ('毛利率', S_M, 'gm', PCT),
            ('期间费用率（销售+管理+研发）', S_OPEX, 'opex_r', PCT), ('归母净利润', S_M, 'np', NUM), ('  同比', S_M, 'np_g', PCT),
            ('EPS（期末流通股本，元）', S_M, 'eps', PS), ('实际税率', S_TAX, 'etr', PCT), ('ROE（杜邦，平均）', S_RET, 'roe', PCT), ('ROIC', S_RET, 'roic', PCT),
            ('资本开支（历史=现金口径；预测=新增投入）', S_CAP, 'capex_tot', NUM), ('折旧摊销', S_CAP, 'da_tot', NUM), ('营运资本/收入', S_WC, 'nwc_r', PCT),
            ('自由现金流（经营现金流-资本开支）', S_M, 'fcf', NUM), ('净现金（含一年以上定期存款）', S_FIN, 'netcash', NUM), ('员工人数（期末）', S_OPEX, 'hc_tot', NUM0), ('PE（当前股价，流通股口径）', S_RET, 'pe', MULT)]
    for lab, sh, k, fmt in outs:
        put(wv, f'A{r}', lab, bold=not lab.startswith('  '))
        for j, mc in enumerate(['G', 'H', 'I', 'J', 'K', 'L']):
            put(wv, f'{CL(3+j)}{r}', f"={X(sh, k, mc)}", fmt)
        r += 1
    r += 1
    section(wv, r, '估值摘要', 9); r += 1
    summ = [('当前情景（1基准/2乐观/3悲观）', f"={A1('scen')}", '0'), ('A股股价（元）', f"={A1('price')}", PS),
            ('WACC', f"={DVX('wacc')}", PCT), ('永续增长率', f"={DVX('g')}", PCT),
            ('DCF每股价值（元，2026-12-31）', f"={DVX('vps')}", PS), ('DCF每股价值折回当前（元）', f"={DVX('vps_now')}", PS),
            ('相对当前股价空间', f"={DVX('up')}", PCT), ('DCF隐含2027E PE', f"={DVX('pe_imp')}", MULT),
            ('反向DCF：当前股价隐含永续增长率', f"={DVX('g_imp')}", PCT2), ('当前股价对应2027E PE', f"={DVX('mkt_pe27')}", MULT),
            ('PE区间估值（2027E EPS × PE下限/上限，元）', f"=TEXT({A1('pe_lo')}*{MM('eps','I')},\"0.0\")&\" - \"&TEXT({A1('pe_hi')}*{MM('eps','I')},\"0.0\")", None),
            ('模型校验汇总（应为0）', f"='校验'!C{REG[('校验','sum')]}", NUM)]
    for lab, f, fmt in summ:
        put(wv, f'A{r}', lab); put(wv, f'C{r}', f, fmt, bold=True)
        REG[('封面', lab)] = r
        r += 1
    r += 1
    section(wv, r, '工作表地图', 9); r += 1
    MAP = [('假设', '情景三组值、敏感性冲击、逐年驱动（带历史参照）'), ('收入拆解', '公司量价、电子元件与通信器件量价拆分、SOFC隔膜片子模型、收入增量与地区'),
           ('成本拆解', '分部毛利率与营业成本；营业成本按原材料/人工/折旧/其他制造费用拆分'), ('人员与费用', '五类员工人数、人均薪酬、职工薪酬按功能归集；销售/管理/研发费用按性质'),
           ('资本开支与折旧', '维护与扩产资本开支、H股境外项目、在建工程滚动、四类固定资产原值/折旧/净值、无形资产、折旧分摊'),
           ('营运资金', '应收票据/账款/融资、存货三分类与跌价准备、应付票据/账款及其他经营负债、周转天数与现金转换周期'),
           ('融资与资金', '受限资金、理财、定期存款、利息收入与理财收益、借款与利息、所有者权益变动、分红回购、净现金'),
           ('税务与其他损益', '税金及附加、其他收益三分项、减值、所得税桥（法定税率→税率差异→加计扣除→其他）'),
           ('股本与每股', 'H股发行、回购、库存股、加权股本、分红'), ('三表预测', '利润表、资产负债表、现金流量表（间接法）联动'),
           ('回报分析', '杜邦ROE、NOPAT、ROIC与增量ROIC、现金转换、估值倍数'), ('DCF估值', '三阶段FCFF、反向DCF、两组实时敏感性表'),
           ('敏感性分析', '龙卷风（2027E归母净利润与DCF每股价值）、二维敏感性、隐含WACC与第二阶段增速求解'), ('情景对比', '三情景关键输出'),
           ('披露明细', '年报/半年报附注原表（固定资产、在建工程、员工、成本构成、存货、税率调节等）')]
    for nm, d in MAP:
        put(wv, f'A{r}', nm, bold=True); put(wv, f'C{r}', d); r += 1
    r += 1
    put(wv, f'A{r}', '免责声明：本模型基于公开信息与产业链调研整理，假设具有主观性，不构成任何投资建议。', italic=True, color='595959')
    wv.column_dimensions['A'].width = 44; wv.column_dimensions['B'].width = 4
    for j in range(6):
        wv.column_dimensions[CL(3 + j)].width = 12
    wv.sheet_view.showGridLines = False

    from summary_v2 import build_summary
    build_summary(wb, g)

    # ================================================================ order
    order = ['封面与摘要', '财务汇总', 'Summary_EN', '假设', S_REV, S_COST, S_OPEX, S_CAP, S_WC, S_FIN, S_TAX, S_SH, S_M, S_RET, 'DCF估值', '相对估值', '敏感性分析', '情景对比',
             'SOFC敏感性', '一致预期对比', '季度数据', '历史利润表', '历史资产负债表', '历史现金流量表', '披露明细', '调研要点', '校验']
    wb._sheets = [wb[n] for n in order]
    for w in wb.worksheets:
        w.sheet_view.showGridLines = False
