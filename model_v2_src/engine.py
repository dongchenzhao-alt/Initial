# -*- coding: utf-8 -*-
"""Two-pass sheet layout engine for the 三环集团 model (rows registered first, formulas written second)."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI
from openpyxl.comments import Comment

F = 'Arial'
BLUE, BLACK, GREEN = '0000FF', '000000', '008000'
NUM = '#,##0.0;(#,##0.0);"-"'
NUM0 = '#,##0;(#,##0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
PCT2 = '0.00%;(0.00%);"-"'
MULT = '0.0"x"'
PS = '0.00;(0.00);"-"'
DAYS = '0;(0);"-"'
HDR_FILL = PatternFill('solid', fgColor='1F3864')
SEC_FILL = PatternFill('solid', fgColor='D9E1F2')
KEY_FILL = PatternFill('solid', fgColor='FFFF00')
FC_FILL = PatternFill('solid', fgColor='F2F2F2')
SUB_FILL = PatternFill('solid', fgColor='EDEDED')

HIST = ['2021A', '2022A', '2023A', '2024A', '2025A']
FCST = ['2026E', '2027E', '2028E', '2029E', '2030E']
YEARS = HIST + FCST
YC = {y: CL(3 + i) for i, y in enumerate(YEARS)}      # C..L
FCOL = [YC[y] for y in FCST]
FY = {'2021A': 'FY2021', '2022A': 'FY2022', '2023A': 'FY2023', '2024A': 'FY2024', '2025A': 'FY2025'}

REG = {}          # (sheet, key) -> row
CUR = [None]


def prev(col):
    return CL(CI(col) - 1)


def put(ws, ref, v, fmt=None, bold=False, color=None, fill=None, italic=False, align=None, size=10, wrap=False):
    c = ws[ref]
    c.value = v
    if color is None:
        if isinstance(v, str) and v.startswith('='):
            color = GREEN if '!' in v else BLACK
        elif isinstance(v, (int, float)):
            color = BLUE
        else:
            color = BLACK
    c.font = Font(name=F, size=size, bold=bold, color=color, italic=italic)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align or wrap:
        c.alignment = Alignment(horizontal=align, vertical='top' if wrap else 'center', wrap_text=wrap)
    return c


def note(ws, ref, text):
    ws[ref].comment = Comment(text, '模型')


def header_row(ws, row, labels, start_col=1, fill=HDR_FILL):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=lab)
        c.font = Font(name=F, size=10, bold=True, color='FFFFFF')
        c.fill = fill
        c.alignment = Alignment(horizontal='center' if i else 'left', vertical='center')


def section(ws, row, text, ncol=13):
    for col in range(1, ncol + 1):
        ws.cell(row=row, column=col).fill = SEC_FILL
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=F, size=10, bold=True, color='1F3864')


def title(ws, text, sub):
    put(ws, 'A1', text, bold=True, size=14, color='1F3864')
    put(ws, 'A2', sub, italic=True, color='595959')


def X(sheet, key, col, absolute=False):
    r = REG[(sheet, key)]
    if absolute:
        ref = f'${col}${r}'
    else:
        ref = f'{col}{r}'
    return ref if sheet == CUR[0] else f"'{sheet}'!{ref}"


class Sheet:
    """Year-column schedule sheet: A=label, B=unit/notes, C..L = 2021A..2030E, M = source/notes."""

    def __init__(self, name, ttl, sub, start=5, header=True, extra_cols=None):
        self.name, self.ttl, self.sub, self.start, self.header = name, ttl, sub, start, header
        self.items = []
        self.extra_cols = extra_cols or []

    def sec(self, text):
        self.items.append(('sec', text))

    def blank(self):
        self.items.append(('blank',))

    def hdr(self, labels):
        self.items.append(('hdr', labels))

    def row(self, key, label, unit='', hist=None, fcst=None, fmt=NUM, bold=False, note=None, src=None, total=False):
        self.items.append(('row', dict(key=key, label=label, unit=unit, hist=hist, fcst=fcst, fmt=fmt, bold=bold,
                                       note=note, src=src, total=total)))

    def layout(self):
        r = self.start
        for it in self.items:
            if it[0] == 'row' and it[1]['key']:
                REG[(self.name, it[1]['key'])] = r
            r += 1

    def write(self, wb):
        ws = wb.create_sheet(self.name)
        CUR[0] = self.name
        title(ws, self.ttl, self.sub)
        if self.header:
            header_row(ws, self.start - 1, ['项目（百万元，另注明除外）', '单位'] + YEARS + ['来源/说明'])
        r = self.start
        for it in self.items:
            kind = it[0]
            if kind == 'sec':
                section(ws, r, it[1])
            elif kind == 'hdr':
                header_row(ws, r, it[1], fill=PatternFill('solid', fgColor='44546A'))
            elif kind == 'row':
                d = it[1]
                put(ws, f'A{r}', d['label'], bold=d['bold'])
                put(ws, f'B{r}', d['unit'], color='595959')
                h = d['hist']
                if h is not None:
                    for y in HIST:
                        v = h(y) if callable(h) else h.get(y)
                        if v is not None:
                            put(ws, f'{YC[y]}{r}', v, d['fmt'], bold=d['bold'])
                f = d['fcst']
                if f is not None:
                    for i, col in enumerate(FCOL):
                        v = f[i] if isinstance(f, (list, tuple)) else f(col)
                        if v is not None:
                            put(ws, f'{col}{r}', v, d['fmt'], bold=d['bold'], fill=FC_FILL)
                if d['src']:
                    put(ws, f'M{r}', d['src'], italic=True, color='595959')
                if d['note']:
                    note(ws, f'A{r}', d['note'])
            r += 1
        ws.column_dimensions['A'].width = 50
        ws.column_dimensions['B'].width = 9
        for col in YC.values():
            ws.column_dimensions[col].width = 11.5
        ws.column_dimensions['M'].width = 70
        ws.freeze_panes = f'C{self.start}'
        ws.sheet_view.showGridLines = False
        return ws
