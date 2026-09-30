# -*- coding: utf-8 -*-
"""Render sens_results.json (static what-if results) into '敏感性分析' and '情景对比'."""
import json, os
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import PatternFill
from engine import *

HERE = os.path.dirname(os.path.abspath(__file__))
SUBHDR = PatternFill('solid', fgColor='44546A')
STATIC = '595959'   # grey = static engine output


def render_sens(wb, g, DV):
    p = os.path.join(HERE, 'sens_results.json')
    ws = wb.create_sheet('敏感性分析'); CUR[0] = '敏感性分析'
    wsc = wb.create_sheet('情景对比')
    if not os.path.exists(p):
        title(ws, '敏感性分析', '尚未运行重算引擎（sens_v2.py）。')
        title(wsc, '情景对比', '尚未运行重算引擎（sens_v2.py）。')
        return
    d = json.load(open(p))
    base = d['base']
    title(ws, '敏感性分析：龙卷风、二维敏感性与反向DCF求解（静态结果）',
          '灰色数字为静态结果：由重算引擎在LibreOffice中逐项改变“假设”第四节冲击单元（或估值参数）后整表重算得到，其他假设保持基准情景不变。修改假设后不会自动更新（“校验”表有提示），需重跑sens_v2.py；DCF表内两张敏感性表与反向DCF为实时公式。')
    r = 4
    section(ws, r, 'A. 基准值（基准情景，全部冲击为0）', 13); r += 1
    for lab, k, fmt in (('2026E归母净利润', 'np26', NUM), ('2027E归母净利润', 'np27', NUM), ('2030E归母净利润', 'np30', NUM),
                        ('2027E自由现金流', 'fcf27', NUM), ('DCF每股价值（元，2026-12-31）', 'vps', PS), ('DCF每股价值折回当前（元）', 'vps_now', PS),
                        ('当前A股股价（元）', None, PS), ('相对当前股价空间', 'up', PCT)):
        put(ws, f'A{r}', lab)
        put(ws, f'C{r}', d['price'] if k is None else base[k], fmt, color=STATIC)
        if k:
            REG[('敏感性分析', k)] = r
        r += 1
    r += 1

    def tornado(r, ttl, items, sort_key, chart_key, chart_title):
        section(ws, r, ttl, 13); r += 1
        header_row(ws, r, ['变量（冲击幅度）', '不利情形输入', '有利情形输入', '2027E归母净利润·不利', '2027E归母净利润·有利', 'Δ净利润·不利', 'Δ净利润·有利',
                           'DCF每股价值·不利', 'DCF每股价值·有利', 'Δ每股价值·不利', 'Δ每股价值·有利', '区间宽度']); r += 1
        items = sorted(items, key=sort_key, reverse=True)
        r0 = r
        for t in items:
            fmt_in = '0' if t['kind'] == '天' else ('0.00' if t['key'] == 'beta' else PCT)
            put(ws, f'A{r}', t['label'])
            put(ws, f'B{r}', t['lo_in'], fmt_in, color=STATIC); put(ws, f'C{r}', t['hi_in'], fmt_in, color=STATIC)
            put(ws, f'D{r}', t['lo']['np27'], NUM, color=STATIC); put(ws, f'E{r}', t['hi']['np27'], NUM, color=STATIC)
            put(ws, f'F{r}', t['lo']['np27'] - base['np27'], NUM, color=STATIC); put(ws, f'G{r}', t['hi']['np27'] - base['np27'], NUM, color=STATIC)
            put(ws, f'H{r}', t['lo']['vps'], PS, color=STATIC); put(ws, f'I{r}', t['hi']['vps'], PS, color=STATIC)
            put(ws, f'J{r}', t['lo']['vps'] - base['vps'], PS, color=STATIC); put(ws, f'K{r}', t['hi']['vps'] - base['vps'], PS, color=STATIC)
            put(ws, f'L{r}', abs(t['hi']['vps'] - t['lo']['vps']) if chart_key == 'vps' else abs(t['hi']['np27'] - t['lo']['np27']),
                PS if chart_key == 'vps' else NUM, color=STATIC)
            r += 1
        ch = BarChart(); ch.type = 'bar'; ch.grouping = 'clustered'; ch.overlap = 100; ch.gapWidth = 40
        ch.title = chart_title; ch.height = max(8, 0.75 * len(items)); ch.width = 20
        cols = ('J', 'K') if chart_key == 'vps' else ('F', 'G')
        for col in cols:
            ci = ord(col) - 64
            ch.add_data(Reference(ws, min_col=ci, min_row=r0 - 1, max_row=r - 1), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=r0, max_row=r - 1))
        ch.y_axis.title = '较基准变动'; ch.x_axis.scaling.orientation = 'maxMin'
        ch.x_axis.tickLblPos = 'low'; ch.x_axis.delete = False; ch.y_axis.delete = False
        ch.y_axis.majorGridlines = None
        ch.legend.position = 'b'
        ws.add_chart(ch, f'N{r0 - 1}')
        return max(r + 1, r0 + int(ch.height * 2) + 1)

    r = tornado(r, 'B. 龙卷风：DCF每股价值（按区间宽度排序；含估值参数）', d['tornado'], lambda t: abs(t['hi']['vps'] - t['lo']['vps']), 'vps',
                'DCF每股价值对各假设的敏感性（元/股）')
    op_items = [t for t in d['tornado'] if t['key'] not in ('rf', 'erp', 'beta', 'g', 'g2s', 'ronic')]
    r = tornado(r, 'C. 龙卷风：2027E归母净利润（按区间宽度排序；经营假设）', op_items, lambda t: abs(t['hi']['np27'] - t['lo']['np27']), 'np27',
                '2027E归母净利润对各经营假设的敏感性（百万元）')

    section(ws, r, 'D. 二维敏感性（行×列同时冲击，冲击叠加在基准情景各年采用值之上）', 13); r += 1
    OLAB = {'np27': ('2027E归母净利润（百万元）', NUM), 'vps': ('DCF每股价值（元）', PS), 'fcf27': ('2027E自由现金流（百万元）', NUM)}
    for gr in d['grids']:
        put(ws, f'A{r}', gr['title'], bold=True, color='1F3864'); r += 1
        rfmt = '0' if gr['ru'] == '天' else PCT
        cfmt = '0' if gr['cu'] == '天' else PCT
        outs = list(gr['res'].keys())
        for bi, o in enumerate(outs[:2]):
            c0 = 2 if bi == 0 else 9
            lab, fmt = OLAB[o]
            put(ws, f'{CL(c0)}{r}', lab, bold=True)
        r += 1
        for bi, o in enumerate(outs[:2]):
            c0 = 2 if bi == 0 else 9
            put(ws, f'{CL(c0)}{r}', '行\\列', bold=True)
            for j, cv in enumerate(gr['cvals']):
                put(ws, f'{CL(c0 + 1 + j)}{r}', cv, cfmt, bold=True, color=STATIC, fill=SEC_FILL)
        r += 1
        for i, rv in enumerate(gr['rvals']):
            for bi, o in enumerate(outs[:2]):
                c0 = 2 if bi == 0 else 9
                lab, fmt = OLAB[o]
                put(ws, f'{CL(c0)}{r}', rv, rfmt, bold=True, color=STATIC, fill=SEC_FILL)
                for j, v in enumerate(gr['res'][o][i]):
                    put(ws, f'{CL(c0 + 1 + j)}{r}', v, fmt, color=STATIC, fill=KEY_FILL if (rv == 0 and gr['cvals'][j] == 0) else None)
            r += 1
        r += 1

    section(ws, r, 'E. 反向DCF：让DCF每股价值（折回当前）等于当前股价，需要什么（每次只动一个变量）', 13); r += 1
    header_row(ws, r, ['变量', '基准值', '所需值', '所需变化', '对应2027E归母净利润', '说明']); r += 1
    rev = d['rev']
    rows = [('beta', 'WACC（通过Beta求解）', lambda v: v['out']['wacc'], lambda: base['wacc'], PCT2, 'WACC需降至约{:.1%}（对应Beta约{:.2f}）'),
            ('g2s', '第二阶段起点增速溢价（相对2030E收入增速）', lambda v: v['x'], lambda: rev['g2s']['base'], PCT, '2031年起增速需在2030E收入增速之上再加该幅度，再线性回落（再投资率按增速÷RONIC同步提高）'),
            ('s_ec_vol', '电子元件销量增速（逐年额外增加）', lambda v: v['x'], lambda: 0.0, PCT, '2026E-2030E每年销量增速均需在基准之上再加该幅度'),
            ('s_sofc_gw', 'Bloom电池片层需求（倍数冲击）', lambda v: v['x'], lambda: 0.0, PCT, '需求需为基准的(1+该值)倍，隔膜片份额与单价不变'),
            ('s_gm', '全部分部毛利率（同时提升）', lambda v: v['x'], lambda: 0.0, PCT, '')]
    UNSOLVED = {'g2s': ('>18pp', '溢价加到18pp（第二阶段再投资率接近100%）后DCF价值仍低于股价，合理区间内无解'),
                's_sofc_gw': ('>+200%', 'Bloom需求提高到基准的3倍后DCF价值仍低于股价，合理区间内无解'),
                's_gm': ('>30pp', '全部分部毛利率提高30pp后DCF价值仍低于股价，合理区间内无解'),
                's_ec_vol': ('>60pp', '销量增速每年再加60pp后DCF价值仍低于股价'),
                'beta': ('<0.5', 'Beta降至0.5（WACC约5%）DCF价值仍低于股价')}
    for key, lab, fx, fb, fmt, txt in rows:
        if key == 's_gm' and rev.get(key, {}).get('solved'):
            txt = '全部分部毛利率需同时提高该幅度'
        v = rev.get(key)
        put(ws, f'A{r}', lab)
        put(ws, f'B{r}', fb(), fmt, color=STATIC)
        if v and v.get('solved'):
            x = fx(v)
            put(ws, f'C{r}', x, fmt, color=STATIC, fill=KEY_FILL)
            put(ws, f'D{r}', x - fb(), fmt, color=STATIC)
            put(ws, f'E{r}', v['out']['np27'], NUM, color=STATIC)
            if key == 'beta':
                put(ws, f'F{r}', txt.format(v['out']['wacc'], v['x']), italic=True, color='595959')
            else:
                put(ws, f'F{r}', txt, italic=True, color='595959')
        else:
            u = UNSOLVED.get(key, ('无解', '合理区间内无解'))
            put(ws, f'C{r}', u[0], color=STATIC, fill=KEY_FILL)
            put(ws, f'F{r}', u[1] + '（区间上限时与股价差{:.1f}元）'.format(v['f_hi'] if v else 0), italic=True, color='595959')
        r += 1
    put(ws, f'A{r}', '永续增长率（实时公式，见DCF估值表）')
    put(ws, f'B{r}', f"='DCF估值'!{DV['g']}", PCT2); put(ws, f'C{r}', f"='DCF估值'!{DV['g_imp']}", PCT2, fill=KEY_FILL)
    put(ws, f'D{r}', f"=C{r}-B{r}", PCT2); put(ws, f'F{r}', 'WACC与前两阶段现金流不变时，股价隐含的永续增长率', italic=True, color='595959')
    r += 2
    section(ws, r, 'F. SOFC隔膜片对估值的贡献（将Bloom需求冲击设为-100%，其余不变）', 13); r += 1
    header_row(ws, r, ['指标', '基准', '无SOFC', 'SOFC贡献', '占比']); r += 1
    zs = d['zero_sofc']
    for lab, a, b, fmt in (('2027E归母净利润', base['np27'], zs['np27'], NUM), ('DCF每股价值（元，2026-12-31）', base['vps'], zs['vps'], PS)):
        put(ws, f'A{r}', lab); put(ws, f'B{r}', a, fmt, color=STATIC); put(ws, f'C{r}', b, fmt, color=STATIC); put(ws, f'D{r}', a - b, fmt, color=STATIC)
        put(ws, f'E{r}', (a - b) / a if a else 0, PCT, color=STATIC)
        r += 1
    ws.column_dimensions['A'].width = 40
    for j in range(2, 14):
        ws.column_dimensions[CL(j)].width = 12.5
    ws.sheet_view.showGridLines = False

    # ================================================================ 情景对比
    CUR[0] = '情景对比'
    title(wsc, '情景对比：基准/乐观/悲观', 'A部分为“假设”表中各情景输入值（实时链接）；B、C部分为重算引擎分别切换情景后得到的静态结果。')
    r = 4
    section(wsc, r, 'A. 关键驱动（2026E-2030E，实时链接“假设”）', 12); r += 1
    header_row(wsc, r, ['驱动', '情景', '2026E', '2027E', '2028E', '2029E', '2030E']); r += 1
    DRV = [('ec_vol', '电子元件 销量增速', PCT), ('ec_px', '电子元件 单价变动', PCT), ('cd_vol', '通信器件 销量增速', PCT), ('cd_px', '通信器件 单价变动', PCT),
           ('gm_ec', '电子元件 毛利率', PCT), ('gm_cd', '通信器件 毛利率', PCT), ('sofc_gw', 'Bloom电池片层需求（GW）', '0.00'),
           ('sofc_sh', '三环隔膜片份额', PCT), ('sofc_px', '隔膜片单价（元/片）', '0.00'), ('cap_dom', '国内扩产资本开支', NUM)]
    for key, lab, fmt in DRV:
        for nm, lab2 in (('base', '基准'), ('bull', '乐观'), ('bear', '悲观')):
            rr_ = REG.get(('假设sc', f'{key}|{nm}'))
            if rr_ is None:
                continue
            put(wsc, f'A{r}', lab if nm == 'base' else '', bold=nm == 'base'); put(wsc, f'B{r}', lab2, color='595959')
            for j, col in enumerate(FCOL):
                put(wsc, f'{CL(3 + j)}{r}', f"='假设'!{col}{rr_}", fmt)
            r += 1
    r += 1
    section(wsc, r, 'B. 关键输出（静态）', 12); r += 1
    YRS = ['2025A', '2026E', '2027E', '2028E', '2029E', '2030E']
    header_row(wsc, r, ['指标', '情景'] + YRS); r += 1
    yr_hdr = r - 1
    MET = [('rev', '营业收入', NUM), ('rev_g', '营业收入同比', PCT), ('sofc', 'SOFC隔膜片收入', NUM), ('gm', '毛利率', PCT), ('np', '归母净利润', NUM),
           ('np_g', '归母净利润同比', PCT), ('eps', 'EPS（元）', PS), ('ebitda', 'EBITDA', NUM), ('capex', '资本开支（流出为负）', NUM), ('fcf', '自由现金流', NUM),
           ('roic', 'ROIC', PCT), ('netcash', '净现金', NUM)]
    np_rows = {}
    for k, lab, fmt in MET:
        for sn, lab2 in (('1', '基准'), ('2', '乐观'), ('3', '悲观')):
            put(wsc, f'A{r}', lab if sn == '1' else '', bold=sn == '1'); put(wsc, f'B{r}', lab2, color='595959')
            for j, y in enumerate(YRS):
                v = d['scen'][sn].get(f'{k}|{y}')
                if v is not None:
                    put(wsc, f'{CL(3 + j)}{r}', v, fmt, color=STATIC)
            if k == 'np':
                np_rows[sn] = r
            r += 1
    r += 1
    section(wsc, r, 'C. 估值（静态）', 12); r += 1
    header_row(wsc, r, ['指标', '', '基准', '乐观', '悲观']); r += 1
    for k, lab, fmt in (('vps', 'DCF每股价值（元，2026-12-31）', PS), ('vps_now', 'DCF每股价值折回当前（元）', PS), ('up', '相对当前股价空间', PCT),
                        ('pe_imp', 'DCF隐含2027E PE', MULT), ('ev_ebitda', 'DCF隐含2027E EV/EBITDA', MULT), ('g_imp', '当前股价隐含永续增长率', PCT2)):
        put(wsc, f'A{r}', lab)
        for j, sn in enumerate(('1', '2', '3')):
            put(wsc, f'{CL(3 + j)}{r}', d['scen'][sn][f'{k}|'], fmt, color=STATIC)
        r += 1
    ch = LineChart(); ch.title = '归母净利润：三情景（百万元）'; ch.height = 7.5; ch.width = 16
    for sn in ('1', '2', '3'):
        ch.add_data(Reference(wsc, min_col=3, max_col=8, min_row=np_rows[sn]), from_rows=True, titles_from_data=False)
    from openpyxl.chart.series import SeriesLabel
    for s_, nm in zip(ch.series, ('基准', '乐观', '悲观')):
        s_.tx = SeriesLabel(v=nm); s_.smooth = False
    ch.x_axis.delete = False; ch.y_axis.delete = False; ch.y_axis.numFmt = '#,##0'
    ch.set_categories(Reference(wsc, min_col=3, max_col=8, min_row=yr_hdr))
    ch.legend.position = 'b'
    wsc.add_chart(ch, 'J5')
    wsc.column_dimensions['A'].width = 30; wsc.column_dimensions['B'].width = 8
    for j in range(3, 9):
        wsc.column_dimensions[CL(j)].width = 12
    wsc.sheet_view.showGridLines = False
